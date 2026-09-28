"""Regression scenarios from the quality review; only synthetic data and I/O."""

import json
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace
from urllib.error import URLError, HTTPError
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.utils.cell import column_index_from_string

import src.web.application as web
import src.services.archive_service as archive_module
from src.field_map import FIELD_MAP
from src.services.file_service import FileService
from src.services.file_mutex import file_mutex, FileMutexBusy
from src.services.local_changes import DIRTY_MARKER, mark_dirty
from src.services.editing_state_service import EditingStateService, EditorIdentity, EditingStateError, OperationInProgressError, RevisionConflictError
from src.services.graph_storage_service import GraphStorageService, GraphConfig, GraphStorageError, GraphConflictError
from src.services.graph_mail_service import GraphMailService, GraphMailConfig, GraphMailConfigurationError, GraphMailUnknownResultError
from src.services.graph_sync_queue import GraphSyncQueue
from src.services.teams_notification_service import TeamsNotificationService, TeamsNotificationConfig, TeamsNotificationConfigurationError
from src.services.finalization_service import FinalizationService


def synthetic_workbook(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "LINK"
    book.create_sheet("FS")
    for field in FIELD_MAP:
        sheet.cell(2, column_index_from_string(field.column), field.label)
    sheet["A3"] = "2026_9900"
    book.save(path)
    book.close()
    return path


def test_excel_read_has_bounded_xml_passes_and_releases_the_file(tmp_path, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    source = synthetic_workbook(tmp_path / "synthetic.xlsx")
    original_source = ReadOnlyWorksheet._get_source
    passes = []
    def open_source(worksheet):
        if worksheet.title == "LINK":
            passes.append(worksheet.title)
        return original_source(worksheet)
    monkeypatch.setattr(ReadOnlyWorksheet, "_get_source", open_source)
    assert web.ExcelService(source).read_link()["A"] == "2026_9900"
    # One dimension inspection plus one data stream, independent of field count.
    assert len(passes) <= 2
    renamed = source.with_name("renamed.xlsx")
    source.rename(renamed)
    web.ExcelService(renamed).write_link_from_form({"Cliente nome": "Cliente fictício"})
    assert web.ExcelService(renamed).read_link_as_form_data()["Cliente nome"] == "Cliente fictício"
    renamed.unlink()


def test_first_draft_and_finalization_preserve_other_files_cache(tmp_path, monkeypatch):
    active = tmp_path / "active"
    source = synthetic_workbook(active / "2026_9900.xlsx")
    unrelated = synthetic_workbook(active / "2026_9901.xlsx")
    monkeypatch.setattr(archive_module, "EXCEL_ARQUIVADAS_DIR", tmp_path / "archive")
    files = FileService(active)
    app = web.create_app(file_service=files, editing_state_service=EditingStateService(tmp_path / "editing"))
    client = app.test_client()
    assert client.get("/api/files").status_code == 200
    opened = client.get(f"/api/file/{source.stem}/bootstrap?client_id=qa").get_json()
    def metadata(editing, key):
        return {"document_id": editing["document_id"], "client_id": "qa",
            "lease_token": editing["lease"]["token"], "base_revision": editing["revision"], "idempotency_key": key}
    payload = {"customer_name": "Cliente fictício alterado", "client_not_present": True,
        "technician_records": [{"technician": "Técnico fictício", "start_time": "09:00", "end_time": "10:00", "date": "2026-09-08"}],
        "_edit": metadata(opened["editing"], "draft")}
    reads = []
    original_read = web.ExcelService.read_link_as_form_data
    def read(service):
        reads.append(service.file_path)
        return original_read(service)
    monkeypatch.setattr(web.ExcelService, "read_link_as_form_data", read)
    saved = client.post(f"/api/file/{source.stem}/draft", json=payload)
    assert saved.status_code == 200, saved.get_json()
    draft_name = saved.get_json()["file"]
    refreshed = client.get("/api/files").get_json()
    assert unrelated not in reads and source not in reads
    assert next(file for file in refreshed["files"] if file["name"] == draft_name)["summary"]["customer_name"] == payload["customer_name"]
    reopened = client.get(f"/api/file/{draft_name}/bootstrap?client_id=qa").get_json()
    payload["_edit"] = metadata(reopened["editing"], "send")
    reads.clear()
    assert client.post(f"/api/file/{draft_name}/send", json=payload).status_code == 200
    assert client.get("/api/files").status_code == 200
    assert reads == []


def graph_service():
    return GraphStorageService(GraphConfig("test", "test", "test", "test"))


def mail_service(**kwargs):
    return GraphMailService(GraphMailConfig("test", "test", "test", "sender@example.test", **kwargs))


@pytest.mark.parametrize("name", ["../outside", "..\\outside", "C:\\outside", "/outside", ".", "..", "outside\x00", "bundle\\child", "x:y"])
def test_r1_rejects_non_simple_names(tmp_path, name):
    active = tmp_path / "active"
    active.mkdir()
    synthetic_workbook(tmp_path / "outside.xlsx")
    assert FileService(active).get_file_by_name(name) is None


def test_r1_api_cannot_read_a_sibling_excel(tmp_path):
    active = tmp_path / "active"
    active.mkdir()
    synthetic_workbook(tmp_path / "outside.xlsx")
    app = web.create_app(file_service=FileService(active), editing_state_service=EditingStateService(tmp_path / "editing"))
    assert app.test_client().get("/api/file/..%5Coutside").status_code == 404


def test_r1_accepts_a_contained_unicode_filename(tmp_path):
    source = synthetic_workbook(tmp_path / "Folha João.xlsx")
    assert FileService(tmp_path).get_file_by_name("Folha João") == source


@pytest.mark.parametrize("secret", [None, "", "dev-local-secret", "short"])
def test_r2_refuses_weak_server_session_key(monkeypatch, secret):
    monkeypatch.setenv("FS_ENVIRONMENT", "production")
    if secret is None:
        monkeypatch.delenv("FS_SECRET_KEY", raising=False)
    else:
        monkeypatch.setenv("FS_SECRET_KEY", secret)
    with pytest.raises(ValueError, match="FS_SECRET_KEY"):
        web._load_secret_key()


def test_r2_https_mode_sets_secure_cookie(monkeypatch, tmp_path):
    monkeypatch.setenv("FS_ENVIRONMENT", "production")
    app = web.create_app(file_service=FileService(tmp_path / "active"), editing_state_service=EditingStateService(tmp_path / "editing"))
    assert app.config["SESSION_COOKIE_SECURE"] is True
    assert app.secret_key != "dev-local-secret"


def test_r2_local_fallback_is_not_the_known_key(monkeypatch):
    monkeypatch.delenv("FS_SECRET_KEY", raising=False)
    assert len(web._load_secret_key()) >= 32
    assert web._load_secret_key() != "dev-local-secret"


def test_r3_only_one_of_twenty_simultaneous_operations_is_admitted(tmp_path):
    source = synthetic_workbook(tmp_path / "2026_9900_2026-09-08_QA.xlsx")
    service = EditingStateService(tmp_path / "editing")
    identity = EditorIdentity("qa", "Synthetic")
    editing = service.acquire_lease(source, identity, "device")
    def claim(index):
        try:
            service.claim_operation(document_id=editing["document_id"], kind="draft", idempotency_key=str(index),
                identity=identity, client_id="device", lease_token=editing["lease"]["token"], base_revision=editing["revision"])
            return True
        except OperationInProgressError:
            return False
    with ThreadPoolExecutor(max_workers=8) as executor:
        assert sum(executor.map(claim, range(20))) == 1
    with pytest.raises(OperationInProgressError):
        service.save_autosave(document_id=editing["document_id"], identity=identity, client_id="device",
            lease_token=editing["lease"]["token"], base_revision=editing["revision"], idempotency_key="auto", document={})


def test_r3_process_guard_blocks_another_thread(tmp_path):
    lock = tmp_path / "lock"
    def attempt():
        with pytest.raises(FileMutexBusy):
            with file_mutex(lock):
                pytest.fail("Concurrent writer entered")
    with file_mutex(lock), ThreadPoolExecutor(max_workers=1) as executor:
        executor.submit(attempt).result(timeout=5)
    with file_mutex(lock):
        pass


def test_r3_commit_cannot_accept_a_changed_revision(tmp_path):
    source = synthetic_workbook(tmp_path / "draft.xlsx")
    service = EditingStateService(tmp_path / "editing")
    identity = EditorIdentity("qa", "Synthetic")
    editing = service.acquire_lease(source, identity, "device")
    service.claim_operation(document_id=editing["document_id"], kind="draft", idempotency_key="save",
        identity=identity, client_id="device", lease_token=editing["lease"]["token"], base_revision=editing["revision"])
    with service._state_transaction(editing["document_id"]) as state:
        state["revision"] += 1
    with pytest.raises(RevisionConflictError):
        service.commit_operation(document_id=editing["document_id"], idempotency_key="save", path=source, response={})


class SyntheticGraph:
    def __init__(self): self.effects = []
    def active_source_name(self, path): return path.parent.name
    def active_source_etag(self, path): return "test-etag"
    def sync_active_files(self, path): return list(path.glob("*/*.xlsx"))
    def upload_archive_bundle(self, path):
        self.effects.append("upload")
        return [path.name]
    def remove_active_name(self, *args, **kwargs):
        self.effects.append("remove")
        return True


@pytest.fixture
def finalizing_app(tmp_path, monkeypatch):
    active = tmp_path / "active"
    name = "2026_9900_2026-09-08_QA"
    source = synthetic_workbook(active / name / (name + ".xlsx"))
    monkeypatch.setattr(archive_module, "EXCEL_ARQUIVADAS_DIR", tmp_path / "archive")
    graph = SyntheticGraph()
    queue = GraphSyncQueue(graph, tmp_path / "queue.sqlite3", auto_start=False)
    editing_service = EditingStateService(tmp_path / "editing")
    app = web.create_app(file_service=FileService(active), graph_service=graph, graph_sync_queue=queue, editing_state_service=editing_service)
    client = app.test_client()
    editing = client.post(f"/api/file/{name}/lease", json={"client_id": "qa"}).get_json()["editing"]
    payload = {"customer_name": "Cliente fictício", "client_not_present": True,
        "technician_records": [{"technician": "Técnico fictício", "start_time": "09:00", "end_time": "10:00", "date": "2026-09-08"}], "_edit": {
        "document_id": editing["document_id"], "client_id": "qa", "lease_token": editing["lease"]["token"],
        "base_revision": editing["revision"], "idempotency_key": "finalize"}}
    return client, source, queue, graph, editing_service, payload


@pytest.mark.parametrize("failure", ["excel", "json", "signature", "photos", "html", "publish", "queue", "commit"])
def test_r4_each_failed_precommit_phase_preserves_editable_source(finalizing_app, failure):
    client, source, queue, graph, editing, payload = finalizing_app
    original_bytes = source.read_bytes()
    targets = {
        "excel": (web.ExcelService, "write_link_from_form"), "json": (web.DocumentDataService, "write"),
        "signature": (web.SignatureService, "save_from_form_data"), "photos": (web.PhotoAttachmentService, "apply"),
        "html": (web.DocumentArtifactService, "write_html"), "publish": (FinalizationService, "publish"),
        "queue": (queue, "enqueue"), "commit": (editing, "commit_operation"),
    }
    with patch.object(*targets[failure], side_effect=OSError("SYNTHETIC_DISK_FAILURE")):
        response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 500, response.get_json()
    assert source.read_bytes() == original_bytes
    assert client.get(f"/api/file/{source.stem}/bootstrap?client_id=qa").status_code == 200
    queue.run_until_idle()
    assert graph.effects == []


@pytest.mark.parametrize("missing_field", ["technician", "start_time", "end_time", "date", "customer_signer_name", "customer_signature_date"])
def test_finalization_rejects_missing_fields_before_writing_or_queuing(finalizing_app, missing_field):
    client, source, queue, graph, editing, payload = finalizing_app
    original_bytes = source.read_bytes()
    if missing_field.startswith("customer_"):
        payload.update(client_not_present=False, customer_signer_name="Maria Santos", customer_signature_date="2026-09-08")
        payload[missing_field] = ""
    else:
        payload["technician_records"][0][missing_field] = ""
    response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 400
    assert response.get_json()["missing_fields"]
    assert source.read_bytes() == original_bytes
    assert not (source.parent / f"{source.stem}__documento.json").exists()
    queue.run_until_idle()
    assert graph.effects == []


def test_draft_still_accepts_missing_finalization_fields(finalizing_app):
    client, source, queue, graph, editing, payload = finalizing_app
    payload.update(technician_records=[], client_not_present=False)
    response = client.post(f"/api/file/{source.stem}/draft", json=payload)
    assert response.status_code == 200, response.get_json()
    assert response.get_json()["success"] is True


@pytest.mark.parametrize("key,value,label", [
    ("customer_signer_name", "Maria", "Primeiro e último nome"),
    ("customer_signature_date", "2026-02-30", "Data da assinatura"),
])
def test_finalization_rejects_invalid_signature_details_without_side_effects(finalizing_app, key, value, label):
    client, source, queue, graph, editing, payload = finalizing_app
    original_bytes = source.read_bytes()
    payload.update(client_not_present=False, customer_signer_name="Maria Santos", customer_signature_date="2026-09-08")
    payload[key] = value
    response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 400
    assert response.get_json()["invalid_fields"] == [label]
    assert source.read_bytes() == original_bytes
    queue.run_until_idle()
    assert graph.effects == []


def test_r4_queue_waits_for_commit_then_publishes_once(finalizing_app):
    client, source, queue, graph, editing, payload = finalizing_app
    response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 200, response.get_json()
    job = queue.status(response.get_json()["graph_job_id"])
    assert job["status"] == "held"
    assert not source.exists()
    assert graph.effects == []
    queue.run_until_idle()
    queue.run_until_idle()
    assert graph.effects == ["upload", "remove"]
    assert queue.status(job["id"])["status"] == "complete"


def test_r4_completed_archive_recovers_interrupted_source_cleanup(finalizing_app):
    client, source, queue, graph, editing, payload = finalizing_app
    with patch.object(FinalizationService, "retire_source", side_effect=OSError("SYNTHETIC_CLEANUP_FAILURE")):
        assert client.post(f"/api/file/{source.stem}/send", json=payload).status_code == 200
    assert source.exists()
    assert client.get("/api/files").status_code == 200
    assert not source.exists()


def test_r4_failed_commit_can_resume_the_same_archive_without_duplicate_publication(finalizing_app):
    client, source, queue, graph, editing, payload = finalizing_app
    with patch.object(editing, "commit_operation", side_effect=OSError("SYNTHETIC_COMMIT_FAILURE")):
        assert client.post(f"/api/file/{source.stem}/send", json=payload).status_code == 500
    queue.run_until_idle()
    assert graph.effects == [] and source.exists()
    response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 200, response.get_json()
    assert len(list((source.parent.parent.parent / "archive").glob("*/*.xlsx"))) == 1
    queue.run_until_idle()
    assert graph.effects == ["upload", "remove"]


def test_r4_finalization_cannot_overtake_an_unfinished_draft_upload(finalizing_app):
    client, source, queue, graph, editing, payload = finalizing_app
    queue.enqueue("upload_active", {"draft_path": str(source)})
    response = client.post(f"/api/file/{source.stem}/send", json=payload)
    assert response.status_code == 409
    assert response.get_json()["code"] == "publication_pending"
    assert source.exists() and graph.effects == []


def test_r4_finalization_refuses_a_source_without_remote_version(finalizing_app, monkeypatch):
    client, source, queue, graph, editing, payload = finalizing_app
    monkeypatch.setattr(graph, "active_source_etag", lambda path: None)
    assert client.post(f"/api/file/{source.stem}/send", json=payload).status_code == 409
    assert source.exists() and graph.effects == []


def test_r4_completed_document_cannot_be_finalized_with_another_key(tmp_path):
    source = synthetic_workbook(tmp_path / "draft.xlsx")
    service = EditingStateService(tmp_path / "editing")
    identity = EditorIdentity("qa", "Synthetic")
    editing = service.acquire_lease(source, identity, "device")
    fields = dict(document_id=editing["document_id"], kind="send", identity=identity,
        client_id="device", lease_token=editing["lease"]["token"], base_revision=editing["revision"])
    service.claim_operation(idempotency_key="send", **fields)
    result = service.commit_operation(document_id=editing["document_id"], idempotency_key="send", path=source, response={})
    fields["base_revision"] = result["revision"]
    with pytest.raises(EditingStateError, match="finalizada"):
        service.claim_operation(idempotency_key="duplicate", **fields)


def test_r5_name_collision_is_terminal_and_never_disables_creation_guard(tmp_path, monkeypatch):
    source = synthetic_workbook(tmp_path / "draft" / "draft.xlsx")
    graph = graph_service()
    monkeypatch.setattr(graph, "_get_item_by_path", lambda path: {"id": "somebody-elses-folder"})
    monkeypatch.setattr(graph, "upload_file", lambda *args, **kwargs: pytest.fail("Must not overwrite"))
    queue = GraphSyncQueue(graph, tmp_path / "queue.sqlite3", auto_start=False)
    job = queue.enqueue("upload_active", {"draft_path": str(source), "fail_if_exists": True})
    queue.run_until_idle()
    first = queue.status(job["id"])
    assert first["status"] == "failed" and first["will_retry"] is False
    queue.retry(job["id"])
    queue.run_until_idle()
    assert queue.status(job["id"])["status"] == "failed"


def test_r5_own_partial_upload_can_resume_without_overwriting_another_folder(tmp_path, monkeypatch):
    source = synthetic_workbook(tmp_path / "draft" / "draft.xlsx")
    graph = graph_service()
    remote = {}
    monkeypatch.setattr(graph, "_get_item_by_path", lambda path: dict(remote) or None)
    monkeypatch.setattr(graph, "ensure_folder_path", lambda path: None)
    def create(method, path, **kwargs):
        assert json.loads(kwargs["data"])["@microsoft.graph.conflictBehavior"] == "fail"
        remote.update(id="our-folder")
        return dict(remote)
    monkeypatch.setattr(graph, "_graph_json", create)
    monkeypatch.setattr(graph, "_upload_photo_folder", lambda *args, **kwargs: [])
    with patch.object(graph, "upload_file", side_effect=GraphStorageError("SYNTHETIC interruption")):
        with pytest.raises(GraphStorageError):
            graph.upload_active_bundle(source, fail_if_exists=True)
    monkeypatch.setattr(graph, "upload_file", lambda file, path, **kwargs: {"id": "file", "eTag": "new"})
    assert graph.upload_active_bundle(source, fail_if_exists=True)


def test_r6_active_listing_reads_all_250_items(monkeypatch):
    graph = graph_service()
    first = [{"id": str(i), "name": f"{i}.xlsx"} for i in range(200)]
    last = [{"id": str(i), "name": f"{i}.xlsx"} for i in range(200, 250)]
    monkeypatch.setattr(graph, "_graph_json", lambda *args: {"value": first, "@odata.nextLink": "https://graph.microsoft.com/v1.0/next"})
    monkeypatch.setattr(graph, "_request_json", lambda *args, **kwargs: {"value": last})
    assert len(graph.list_active_items()) == 250
    assert len(graph._list_children("folder")) == 250


@pytest.mark.parametrize("next_link", ["https://example.test/leak", "http://graph.microsoft.com/v1.0/next", "https://graph.microsoft.com.evil.test/v1.0/next"])
def test_r6_rejects_unsafe_next_page_without_contacting_it(monkeypatch, next_link):
    graph = graph_service()
    monkeypatch.setattr(graph, "_graph_json", lambda *args: {"value": [], "@odata.nextLink": next_link})
    monkeypatch.setattr(graph, "_request_json", lambda *args, **kwargs: pytest.fail("Unsafe pagination request"))
    with pytest.raises(GraphStorageError):
        graph.list_active_items()


def test_r6_failed_pagination_does_not_clean_local_cache(tmp_path, monkeypatch):
    graph = graph_service()
    cached = tmp_path / "cached.xlsx"
    cached.write_bytes(b"SYNTHETIC")
    cached.with_name(cached.name + ".graph.json").write_text("{}")
    monkeypatch.setattr(graph, "_graph_json", lambda *args: {"value": [], "@odata.nextLink": "https://graph.microsoft.com/v1.0/next"})
    monkeypatch.setattr(graph, "_request_json", lambda *args, **kwargs: (_ for _ in ()).throw(GraphStorageError("SYNTHETIC")))
    with pytest.raises(GraphStorageError):
        graph.sync_active_files(tmp_path)
    assert cached.read_bytes() == b"SYNTHETIC"


def test_r7_refresh_preserves_an_unpublished_draft(tmp_path, monkeypatch):
    name = "2026_9900_2026-09-08_QA"
    source = synthetic_workbook(tmp_path / "active" / name / (name + ".xlsx"))
    graph = graph_service()
    queue = GraphSyncQueue(graph, tmp_path / "queue.sqlite3", auto_start=False)
    job = queue.enqueue("upload_active", {"draft_path": str(source)})
    monkeypatch.setattr(graph, "list_active_items", lambda: [])
    graph.sync_active_files(source.parent.parent)
    assert source.exists()
    assert Path(queue.status(job["id"])["payload"]["draft_path"]).exists()


def test_r7_old_upload_cannot_clear_a_newer_local_change(tmp_path):
    source = synthetic_workbook(tmp_path / "draft" / "draft.xlsx")
    queue = GraphSyncQueue(None, tmp_path / "queue.sqlite3", auto_start=False)
    job = queue.enqueue("upload_active", {"draft_path": str(source)})
    new_version = mark_dirty(source.parent)
    queue._publish_staging_metadata(job["payload"])
    assert (source.parent / DIRTY_MARKER).read_text() == new_version


def test_r7_cache_repair_preserves_a_new_directory_before_its_first_write(tmp_path):
    draft = tmp_path / "2026_9900_2026-09-08_QA"
    draft.mkdir()
    graph_service()._repair_active_cache(tmp_path)
    assert draft.is_dir()


def test_r8_lost_response_does_not_duplicate_mail_and_blocks_manual_retry(tmp_path):
    mail = mail_service()
    accepted = []
    def transport(req, **kwargs):
        accepted.append(req.full_url)
        raise URLError("SYNTHETIC: accepted but response lost")
    mail._transport = transport
    mail._access_token = "synthetic"
    mail._access_token_expires_at = time.time() + 3600
    prepared = mail.prepare_service_email(customer_email="qa@example.test", technician_email="qa@example.test", customer_name="Synthetic", service_number="TEST")
    pdf = tmp_path / "test.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%%EOF")
    queue = GraphSyncQueue(None, tmp_path / "queue.sqlite3", mail_service=mail,
        local_pdf_service=SimpleNamespace(export_archive_pdf=lambda path: pdf), auto_start=False)
    job = queue.enqueue("archive_and_mail_local", {"archived_path": "synthetic.xlsx", "mail": prepared})
    queue.run_until_idle()
    queue.run_until_idle()
    result = queue.status(job["id"])
    assert result["reconciliation_required"] is True
    assert result["will_retry"] is False
    with pytest.raises(ValueError, match="reconciliação"):
        queue.retry(job["id"])
    assert len(accepted) == 1


def test_r8_recovery_cannot_erase_unknown_mail_with_a_pdf_checkpoint(tmp_path):
    queue = GraphSyncQueue(None, tmp_path / "queue.sqlite3", auto_start=False)
    with pytest.raises(GraphMailUnknownResultError):
        queue._execute("archive_and_mail_local", {}, prior_result={"phase": "sending_mail"})


@pytest.mark.parametrize("exception", [TimeoutError("synthetic"), HTTPError("https://graph.microsoft.com/v1.0/users/test/sendMail", 503, "synthetic", {}, None)])
def test_r8_send_timeout_and_server_error_are_ambiguous(exception):
    service = mail_service()
    service._transport = lambda *args, **kwargs: (_ for _ in ()).throw(exception)
    with pytest.raises(GraphMailUnknownResultError):
        service._request_bytes("https://graph.microsoft.com/v1.0/users/test/sendMail", method="POST", data=b"{}", authenticated=False)


def test_r9_test_recipient_is_the_only_final_recipient(tmp_path):
    service = mail_service(test_recipient="testbox@example.test")
    prepared = service.prepare_service_email(customer_email="customer@example.test", technician_email="technician@example.test", customer_name="Synthetic", service_number="TEST")
    assert prepared["cc"] == []
    pdf = tmp_path / "test.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%%EOF")
    with patch.object(service, "_graph_bytes", return_value=b"") as mocked:
        result = service.send_prepared(prepared, pdf, operation_id="test")
    message = json.loads(mocked.call_args.kwargs["data"])["message"]
    assert message["toRecipients"] == [{"emailAddress": {"address": "testbox@example.test"}}]
    assert message["ccRecipients"] == [] and not message.get("bccRecipients")
    assert result["cc"] == []


def test_r9_direct_mail_service_cannot_use_a_real_transport_in_test_mode():
    service = mail_service()
    with patch("src.services.graph_mail_service.request.urlopen") as transport:
        with pytest.raises(GraphMailConfigurationError, match="bloqueado"):
            service._request_bytes("https://graph.microsoft.com/v1.0/users/test/sendMail", method="POST", data=b"{}")
    transport.assert_not_called()


def test_r9_teams_transport_is_also_blocked_in_test_mode():
    service = TeamsNotificationService(TeamsNotificationConfig("https://example.test/workflow"))
    prepared = service.prepare_service_sent_notification(service_number="SYNTHETIC")
    with patch("src.services.teams_notification_service.request.urlopen") as transport:
        with pytest.raises(TeamsNotificationConfigurationError, match="bloqueado"):
            service.send_prepared(prepared, operation_id="test")
    transport.assert_not_called()
