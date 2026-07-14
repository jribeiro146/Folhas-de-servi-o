"""
Folhas de Servico - Servico de leitura/escrita Excel.

Este modulo le e escreve os campos na sheet LINK.
Tambem posiciona as assinaturas finais na sheet FS.
Usa o field_map como contrato estavel para acesso aos campos.
"""

from __future__ import annotations

import datetime as dt
import os
import secrets
import time
import gc
import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image as PILImage

from src.config import LINK_DATA_ROW, LINK_LABELS_ROW, REQUIRED_SHEETS, SHEET_LINK, SHEET_TEMPLATE
from src.field_map import FIELD_BY_COLUMN, FIELD_MAP, FieldDef, FieldType
from src.services.signature_service import (
    CLIENT_SIGNATURE_LABEL,
    TECHNICIAN_SIGNATURE_LABEL,
    SignatureService,
)

FS_SIGNATURE_TARGETS = {
    CLIENT_SIGNATURE_LABEL: {"column_start": "V", "column_end": "AR", "row": 63},
}
FS_LEGACY_SIGNATURE_TARGETS = {
    TECHNICIAN_SIGNATURE_LABEL: {"column_start": "B", "column_end": "T", "row": 63},
}
FS_SIGNATURE_ROW_HEIGHT_POINTS = 42.0
FS_SIGNATURE_MARGIN_PIXELS = 6

FIELD_HEADER_ALIASES = {
    "Pedido por": ("Pedido por:",),
    "Data pedido": ("Data\npedido",),
    "Telefone (2)": ("Telefone",),
    "Cód. Postal": ("Cód.\nPostal", "Codigo Postal", "Código Postal"),
    "Relatório Técnico": (
        "Relatório Técnico da Intervenção",
        "Relatório Técnico da Intervenção | Intervention report",
    ),
    "T.Total h (2)": ("T.Total h",),
    "T.Total m (2)": ("T.Total m",),
}

FIELD_HEADER_OCCURRENCES = {
    ("Telefone (2)", "Telefone"): 2,
    ("T.Total h (2)", "T.Total h"): 2,
    ("T.Total m (2)", "T.Total m"): 2,
}

LINK_FORM_INPUT_LABELS = {
    "Folha nº",
    "Pedido por",
    "Email",
    "Telefone",
    "Contacto",
    "Telefone (2)",
    "Contrato nº",
    "Loja nº",
    "Data pedido",
    "Cliente nº",
    "Ident",
    "ASSIST",
    "MAN",
    "COL.SERV",
    "GAR",
    "INST",
    "PIQ",
    "FORM",
    "REP.OF",
    "ACOMP.COM",
    "SADI",
    "CCTV",
    "PA/VA",
    "SAI",
    "EXT",
    "SCA",
    "EAS",
    "SADG",
    "SCH",
    "OTHER",
    "Cliente nome",
    "Local",
    "Morada",
    "Cód. Postal",
    "CP",
    "NIF",
    "Avaria reportada",
}


class ExcelValidationError(Exception):
    """Erro de validacao do ficheiro Excel."""


class ExcelFileLockedError(ExcelValidationError):
    """Erro para ficheiros Excel bloqueados por outro processo."""


class ExcelService:
    """Servico para operacoes de leitura/escrita do workbook."""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(f"Ficheiro Excel nao encontrado: {self.file_path}")

        if self.file_path.suffix.lower() != ".xlsx":
            raise ExcelValidationError(f"Ficheiro nao e .xlsx: {self.file_path.name}")

    def _locked_message(self) -> str:
        return (
            "O ficheiro esta em uso por outro programa. "
            "Feche-o no Excel e aguarde o OneDrive antes de atualizar."
        )

    def _load_workbook(
        self,
        *,
        read_only: bool = False,
        data_only: bool = False,
        retries: int = 4,
        delay_seconds: float = 0.5,
    ):
        last_error: Exception | None = None

        for attempt in range(retries):
            try:
                return openpyxl.load_workbook(
                    str(self.file_path),
                    read_only=read_only,
                    data_only=data_only,
                )
            except PermissionError as exc:
                last_error = exc
                if attempt == retries - 1:
                    raise ExcelFileLockedError(self._locked_message()) from exc
                time.sleep(delay_seconds)

        raise ExcelFileLockedError(self._locked_message()) from last_error

    def validate_sheets(self) -> None:
        wb = self._load_workbook(read_only=True)
        try:
            self._validate_sheet_names(wb.sheetnames)
        finally:
            wb.close()
            gc.collect()

    def read_link(self, fields: list[FieldDef] | tuple[FieldDef, ...] | None = None) -> dict[str, Any]:
        wb = self._load_workbook(read_only=True, data_only=True)
        try:
            self._validate_sheet_names(wb.sheetnames)
            ws = wb[SHEET_LINK]
            field_columns = self._resolve_link_columns(ws)
            data: dict[str, Any] = {}

            for field in fields or FIELD_MAP:
                col_idx = field_columns.get(field.column)
                data[field.column] = ws.cell(row=LINK_DATA_ROW, column=col_idx).value if col_idx else None

            return data
        finally:
            wb.close()
            gc.collect()

    def read_link_as_form_data(self) -> dict[str, Any]:
        input_fields = tuple(field for field in FIELD_MAP if field.label in LINK_FORM_INPUT_LABELS)
        raw_data = self.read_link(input_fields)
        form_data: dict[str, Any] = {}

        for field in input_fields:
            value = raw_data.get(field.column)
            form_data[field.label] = self._format_for_form(field, value)

        return form_data

    def write_link(
        self,
        data: dict[str, Any],
        *,
        signatures: dict[str, str | None] | None = None,
    ) -> None:
        wb = self._load_workbook()
        temp_path = self.file_path.with_name(
            f".{self.file_path.stem}.{secrets.token_hex(6)}.tmp{self.file_path.suffix}"
        )
        try:
            self._validate_sheet_names(wb.sheetnames)
            ws = wb[SHEET_LINK]
            field_columns = self._resolve_link_columns(ws)

            for column, value in data.items():
                field = FIELD_BY_COLUMN.get(column)
                if field is None or field.read_only:
                    continue

                col_idx = field_columns.get(field.column)
                if col_idx is None:
                    continue

                ws.cell(row=LINK_DATA_ROW, column=col_idx).value = value

            if signatures is not None:
                self._write_fs_signatures(wb[SHEET_TEMPLATE], signatures)

            wb.save(str(temp_path))
        except PermissionError as exc:
            raise ExcelFileLockedError(self._locked_message()) from exc
        finally:
            wb.close()
            gc.collect()

        try:
            os.replace(str(temp_path), str(self.file_path))
        except PermissionError as exc:
            raise ExcelFileLockedError(self._locked_message()) from exc
        finally:
            try:
                temp_path.unlink()
            except FileNotFoundError:
                pass

    @staticmethod
    def _validate_sheet_names(sheet_names: list[str]) -> None:
        missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in sheet_names]
        if missing:
            raise ExcelValidationError(
                f"Sheets obrigatorias em falta: {', '.join(missing)}. "
                f"Sheets encontradas: {', '.join(sheet_names)}"
            )

    def write_link_from_form(
        self,
        form_data: dict[str, Any],
        *,
        signatures: dict[str, str | None] | None = None,
    ) -> None:
        column_data: dict[str, Any] = {}

        for field in FIELD_MAP:
            if field.read_only:
                continue
            if field.label in form_data:
                value = form_data[field.label]
                column_data[field.column] = self._format_for_excel(field, value)

        self.write_link(column_data, signatures=signatures)

    def get_folha_numero(self) -> str | None:
        data = self.read_link()
        return data.get("A")

    def validate_required_form_data(self, form_data: dict[str, Any]) -> list[str]:
        missing: list[str] = []

        for field in FIELD_MAP:
            if not field.required or field.read_only:
                continue

            value = form_data.get(field.label)
            if self._is_missing_value(field.field_type, value):
                missing.append(field.label)

        return missing

    @staticmethod
    def _format_for_form(field: FieldDef, value: Any) -> Any:
        if value is None:
            if field.field_type == FieldType.CHECKBOX:
                return False
            if field.field_type in (FieldType.CURRENCY, FieldType.NUMBER):
                return 0
            return ""

        if field.field_type == FieldType.CHECKBOX:
            return str(value).strip().upper() == "X"

        if isinstance(value, dt.datetime):
            return value.strftime("%Y-%m-%d")

        if isinstance(value, dt.date):
            return value.strftime("%Y-%m-%d")

        if isinstance(value, dt.time):
            return value.strftime("%H:%M")

        return value

    @staticmethod
    def _format_for_excel(field: FieldDef, value: Any) -> Any:
        if field.field_type == FieldType.CHECKBOX:
            return "X" if value else None

        if field.field_type in (FieldType.CURRENCY, FieldType.NUMBER):
            if value == "" or value is None:
                return 0
            try:
                return float(value) if "." in str(value) else int(value)
            except (ValueError, TypeError):
                return 0

        if value == "":
            return None

        return value

    @staticmethod
    def _is_missing_value(field_type: FieldType, value: Any) -> bool:
        if value is None:
            return True

        if field_type == FieldType.CHECKBOX:
            return value is not True

        if isinstance(value, str):
            return value.strip() == ""

        return False

    @classmethod
    def _resolve_link_columns(cls, worksheet) -> dict[str, int]:
        header_columns = cls._build_header_columns(worksheet)
        has_link_headers = sum(len(columns) for columns in header_columns.values()) >= 5
        resolved: dict[str, int] = {}

        for field in FIELD_MAP:
            col_idx = cls._find_column_for_field(field, header_columns)
            if col_idx is not None:
                resolved[field.column] = col_idx
                continue

            if not has_link_headers:
                resolved[field.column] = column_index_from_string(field.column)

        return resolved

    @classmethod
    def _build_header_columns(cls, worksheet) -> dict[str, list[int]]:
        header_columns: dict[str, list[int]] = {}

        for col_idx in range(1, worksheet.max_column + 1):
            normalized = cls._normalize_header_label(worksheet.cell(row=LINK_LABELS_ROW, column=col_idx).value)
            if not normalized:
                continue
            header_columns.setdefault(normalized, []).append(col_idx)

        return header_columns

    @classmethod
    def _find_column_for_field(
        cls,
        field: FieldDef,
        header_columns: dict[str, list[int]],
    ) -> int | None:
        for candidate in (field.label, *FIELD_HEADER_ALIASES.get(field.label, ())):
            normalized = cls._normalize_header_label(candidate)
            matches = header_columns.get(normalized, [])
            if not matches:
                continue

            occurrence = FIELD_HEADER_OCCURRENCES.get((field.label, candidate), 1)
            if len(matches) >= occurrence:
                return matches[occurrence - 1]

        return None

    @staticmethod
    def _normalize_header_label(value: Any) -> str:
        text = "" if value is None else str(value)
        text = text.replace("\n", " ").replace("\r", " ")
        text = text.replace("º", "o").replace("ª", "a").replace("°", "o")
        text = unicodedata.normalize("NFKD", text)
        text = "".join(char for char in text if not unicodedata.combining(char))
        text = text.casefold().strip()
        text = re.sub(r"[^a-z0-9]+", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    def _write_fs_signatures(
        self,
        worksheet,
        signatures: dict[str, str | None],
    ) -> None:
        signature_targets = tuple(FS_SIGNATURE_TARGETS.items())
        cleanup_targets = tuple({**FS_LEGACY_SIGNATURE_TARGETS, **FS_SIGNATURE_TARGETS}.items())
        if any(signatures.get(label) for label, _ in signature_targets):
            current_height = worksheet.row_dimensions[63].height or 0
            worksheet.row_dimensions[63].height = max(
                current_height,
                FS_SIGNATURE_ROW_HEIGHT_POINTS,
            )

        for _, target in cleanup_targets:
            self._remove_signature_image(
                worksheet,
                row=target["row"],
                column=target["column_start"],
            )

        for label, target in signature_targets:
            data_url = signatures.get(label)
            if not data_url:
                continue

            image = self._build_signature_image(worksheet, data_url, target)
            image.anchor = f"{target['column_start']}{target['row']}"
            worksheet.add_image(image)

    def _build_signature_image(
        self,
        worksheet,
        data_url: str,
        target: dict[str, str | int],
    ) -> WorksheetImage:
        signature_image = self._load_signature_image(data_url)
        target_width = max(
            1,
            self._estimate_range_width_pixels(
                worksheet,
                target["column_start"],
                target["column_end"],
            ) - (FS_SIGNATURE_MARGIN_PIXELS * 2),
        )
        target_height = max(
            1,
            self._row_height_to_pixels(worksheet.row_dimensions[target["row"]].height)
            - (FS_SIGNATURE_MARGIN_PIXELS * 2),
        )
        scaled_image = self._resize_to_fit(signature_image, target_width, target_height)
        image_buffer = BytesIO()
        scaled_image.save(image_buffer, format="PNG")
        image_buffer.seek(0)
        worksheet_image = WorksheetImage(image_buffer)
        worksheet_image._image_buffer = image_buffer
        return worksheet_image

    @staticmethod
    def _load_signature_image(data_url: str) -> PILImage.Image:
        raw_bytes = SignatureService.decode_data_url(data_url)
        image = PILImage.open(BytesIO(raw_bytes))
        image.load()
        prepared_image = image.convert("RGBA")
        bounds = prepared_image.getbbox()
        if bounds:
            prepared_image = prepared_image.crop(bounds)
        return prepared_image

    @staticmethod
    def _resize_to_fit(
        image: PILImage.Image,
        max_width: int,
        max_height: int,
    ) -> PILImage.Image:
        width, height = image.size
        if width <= 0 or height <= 0:
            return image

        scale = min(max_width / width, max_height / height)
        scaled_width = max(1, int(round(width * scale)))
        scaled_height = max(1, int(round(height * scale)))

        if (scaled_width, scaled_height) == image.size:
            return image

        return image.resize(
            (scaled_width, scaled_height),
            resample=PILImage.Resampling.LANCZOS,
        )

    @staticmethod
    def _remove_signature_image(worksheet, *, row: int, column: str) -> None:
        target_row = row - 1
        target_column = column_index_from_string(column) - 1
        remaining_images = []

        for image in list(getattr(worksheet, "_images", [])):
            if ExcelService._image_anchor_matches(image, target_row, target_column):
                continue
            remaining_images.append(image)

        worksheet._images = remaining_images

    @staticmethod
    def _image_anchor_matches(image, row: int, column: int) -> bool:
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        if marker is not None:
            return marker.row == row and marker.col == column

        if isinstance(anchor, str):
            expected = f"{get_column_letter(column + 1)}{row + 1}"
            return anchor.upper() == expected

        return False

    @staticmethod
    def _estimate_range_width_pixels(worksheet, start_column: str, end_column: str) -> int:
        start_index = column_index_from_string(start_column)
        end_index = column_index_from_string(end_column)
        total_pixels = 0

        for column_index in range(start_index, end_index + 1):
            column_letter = get_column_letter(column_index)
            column_width = worksheet.column_dimensions[column_letter].width
            total_pixels += ExcelService._column_width_to_pixels(column_width)

        return total_pixels

    @staticmethod
    def _column_width_to_pixels(width: float | None) -> int:
        normalized_width = width or 8.43
        if normalized_width <= 1:
            return int(round(normalized_width * 12))
        return int(round((normalized_width * 7) + 5))

    @staticmethod
    def _row_height_to_pixels(height: float | None) -> int:
        normalized_height = height or 15
        return int(round(normalized_height * (4 / 3)))
