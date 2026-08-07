import json
from concurrent.futures import ThreadPoolExecutor
import shutil
import tempfile
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pytest

from src.services.archive_service import ArchiveService
from src.services.editing_state_service import EditingStateService
from src.services.file_service import FileService
from src.services.graph_mail_service import GraphMailConfig, GraphMailService
from src.web.application import create_app

FIXTURE_DIR = Path(__file__).parent / "fixtures"
SAMPLE_FILE = FIXTURE_DIR / "test_sample.xlsx"
SIGNATURE_DATA_URL = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+XgnsAAAAASUVORK5CYII="
)


def copy_sample(destination_dir: Path, name: str = "test_sample.xlsx") -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    target = destination_dir / name
    workbook = openpyxl.load_workbook(SAMPLE_FILE)
    try:
        workbook.save(target)
    finally:
        workbook.close()
    return target


def get_fs_image_positions(workbook_path: Path) -> tuple[set[tuple[int, int]], float | None]:
    workbook = openpyxl.load_workbook(workbook_path)
    try:
        worksheet = workbook["FS"]
        positions: set[tuple[int, int]] = set()

        for image in getattr(worksheet, "_images", []):
            marker = getattr(getattr(image, "anchor", None), "_from", None)
            if marker is None:
                continue
            positions.add((marker.row + 1, marker.col + 1))

        return positions, worksheet.row_dimensions[63].height
    finally:
        workbook.close()


class DummyService:
    pass


def acquire_editing(client, file_name: str, client_id: str = "test-tab") -> dict:
    response = client.post(
        f"/api/file/{file_name}/lease",
        json={"client_id": client_id},
    )
    assert response.status_code == 200, response.get_json()
    return response.get_json()["editing"]


def edit_metadata(
    editing: dict,
    *,
    client_id: str = "test-tab",
    idempotency_key: str | None = None,
) -> dict:
    return {
        "document_id": editing["document_id"],
        "client_id": client_id,
        "lease_token": editing["lease"]["token"],
        "base_revision": editing["revision"],
        "idempotency_key": idempotency_key or str(uuid.uuid4()),
    }


@pytest.fixture
def isolated_dirs(monkeypatch):
    base_dir = Path(tempfile.mkdtemp(prefix=f"folhas-servico-{uuid.uuid4()}-"))
    import src.web.application as application_module

    monkeypatch.setattr(application_module, "ACTIVE_AUTH_PROVIDER", "none")
    monkeypatch.setattr(application_module, "AUTH_ENABLED", False)
    monkeypatch.setattr(application_module, "MAIL_ENABLED", False)
    monkeypatch.setattr(application_module, "TEAMS_NOTIFICATIONS_ENABLED", False)
    monkeypatch.setattr(application_module, "STORAGE_BACKEND", "local")
    active_dir = base_dir / "Excel" / "Activas"
    archived_dir = base_dir / "Excel" / "Arquivadas"
    canceled_dir = base_dir / "Excel" / "Canceladas"

    import src.services.archive_service as archive_module

    monkeypatch.setattr(archive_module, "EXCEL_ARQUIVADAS_DIR", archived_dir)
    monkeypatch.setattr(archive_module, "EXCEL_CANCELADAS_DIR", canceled_dir)
    try:
        yield {
            "active": active_dir,
            "archived": archived_dir,
            "canceled": canceled_dir,
            "editing": base_dir / "editing-state",
        }
    finally:
        shutil.rmtree(base_dir, ignore_errors=True)


def test_archive_service_moves_to_archived(isolated_dirs):
    source = copy_sample(isolated_dirs["active"], "sample.xlsx")
    service = ArchiveService()

    archived = service.archive(source)

    assert archived.exists()
    assert archived.parent == isolated_dirs["archived"] / "sample"
    assert not source.exists()


def test_web_api_flow_send_and_cancel(isolated_dirs):
    active_dir = isolated_dirs["active"]
    copy_sample(active_dir, "2026_4572.xlsx")

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()

    files_response = client.get("/api/files")
    files_payload = files_response.get_json()
    assert files_response.status_code == 200
    assert files_payload["success"] is True
    assert files_payload["files"][0]["name"] == "2026_4572"
    assert files_payload["files"][0]["summary"]["customer_name"] == "Hospital CUF Descobertas, SA"
    assert files_payload["files"][0]["summary"]["local_store"] == "CUF Descobertas"
    assert files_payload["files"][0]["summary"]["contact_label"] == "Pedido por"
    assert files_payload["files"][0]["summary"]["contact_name"] == "Sra. Jéssica Silv"
    assert "R. Mário Botas" in files_payload["files"][0]["summary"]["address"]
    assert files_payload["files"][0]["summary"]["phone"] == "910036247"

    get_response = client.get("/api/file/2026_4572")
    get_payload = get_response.get_json()
    assert get_response.status_code == 200
    assert get_payload["success"] is True
    assert get_payload["data"]["Cliente nome"]

    editing = acquire_editing(client, "2026_4572")
    draft_payload = {
        "customer_name": "Cliente Teste",
        "requested_by": "Pedido Teste",
        "intervention_report": "Teste de rascunho",
        "technician_records": [
            {
                "technician": "João Freire",
                "start_time": "09:00",
                "end_time": "11:30",
                "date": "2026-04-17",
            }
        ],
        "materials": [
            {
                "ref": "MAT-01",
                "description": "Bateria",
                "qty": "1",
            }
        ],
        "Assinatura Cliente": SIGNATURE_DATA_URL,
        "_edit": edit_metadata(editing),
    }
    draft_response = client.post("/api/file/2026_4572/draft", json=draft_payload)
    draft_result = draft_response.get_json()
    assert draft_response.status_code == 200
    assert draft_result["success"] is True
    assert draft_result["status"] == "in_progress"
    assert draft_result["created_copy"] is True
    assert draft_result["file"].startswith("2026_4572_")
    assert draft_result["file"].endswith("_JF")
    assert (isolated_dirs["active"] / "2026_4572.xlsx").exists()
    draft_dir = isolated_dirs["active"] / draft_result["file"]
    draft_excel = draft_dir / f"{draft_result['file']}.xlsx"
    assert draft_excel.exists()
    assert (draft_dir / f"{draft_result['file']}__documento.json").exists()
    assert (draft_dir / f"{draft_result['file']}__assinatura_cliente.png").exists()
    draft_positions, draft_row_height = get_fs_image_positions(draft_excel)
    assert (63, 22) in draft_positions
    assert draft_row_height >= 42

    draft_editing = acquire_editing(client, draft_result["file"], "draft-tab")
    send_idempotency_key = str(uuid.uuid4())
    send_body = {
        "customer_name": "Cliente Final",
        "customer_signer_name": "Maria Santos",
        "customer_signature_date": "2026-04-17",
        "intervention_report": "Fecho",
        "_internal_observations": "Nota interna para consulta da equipa.",
        "technician_records": [
            {
                "technician": "João Freire",
                "date": "2026-04-17",
            }
        ],
        "_edit": edit_metadata(
            draft_editing,
            client_id="draft-tab",
            idempotency_key=send_idempotency_key,
        ),
    }
    send_response = client.post(
        f"/api/file/{draft_result['file']}/send",
        json=send_body,
    )
    send_payload = send_response.get_json()
    assert send_response.status_code == 200
    assert send_payload["success"] is True
    replay_response = client.post(
        f"/api/file/{draft_result['file']}/send",
        json=send_body,
    )
    assert replay_response.status_code == 200
    assert replay_response.get_json()["archived_excel"] == send_payload["archived_excel"]
    assert (isolated_dirs["active"] / "2026_4572.xlsx").exists()
    archived_dir = isolated_dirs["archived"] / draft_result["file"]
    assert (archived_dir / f"{draft_result['file']}.xlsx").exists()
    archived_html = archived_dir / f"{draft_result['file']}__folha_final.html"
    internal_notes = archived_dir / f"{draft_result['file']}__observacoes_internas.txt"
    document_json = archived_dir / f"{draft_result['file']}__documento.json"
    assert archived_html.exists()
    assert document_json.exists()
    assert internal_notes.exists()
    assert internal_notes.read_text(encoding="utf-8").strip() == "Nota interna para consulta da equipa."
    assert (archived_dir / f"{draft_result['file']}__assinatura_cliente.png").exists()
    html_content = archived_html.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in html_content
    assert "Cliente Final" in html_content
    assert "Primeiro e último nome" in html_content
    assert "Maria Santos" in html_content
    assert "2026-04-17" in html_content
    assert "data:image/png;base64," in html_content
    assert "Nota interna para consulta da equipa." not in html_content
    assert "_internal_observations" not in json.loads(document_json.read_text(encoding="utf-8"))
    archived_positions, archived_row_height = get_fs_image_positions(
        archived_dir / f"{draft_result['file']}.xlsx"
    )
    assert (63, 22) in archived_positions
    assert archived_row_height >= 42

    copy_sample(active_dir, "2026_9999_2026-07-14_TT.xlsx")
    cancel_editing = acquire_editing(client, "2026_9999_2026-07-14_TT", "cancel-tab")
    cancel_response = client.post(
        "/api/file/2026_9999_2026-07-14_TT/cancel",
        json={"_edit": edit_metadata(cancel_editing, client_id="cancel-tab")},
    )
    cancel_payload = cancel_response.get_json()
    assert cancel_response.status_code == 200
    assert cancel_payload["success"] is True
    assert (isolated_dirs["canceled"] / "2026_9999_2026-07-14_TT" / "2026_9999_2026-07-14_TT.xlsx").exists()


def test_web_api_rejects_missing_required_fields(isolated_dirs):
    active_dir = isolated_dirs["active"]
    draft_name = "2026_4572_2026-07-14_JF"
    copy_sample(active_dir, f"{draft_name}.xlsx")

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()

    editing = acquire_editing(client, draft_name)
    response = client.post(
        f"/api/file/{draft_name}/send",
        json={
            "intervention_report": "Sem cliente",
            "_edit": edit_metadata(editing),
        },
    )
    payload = response.get_json()

    assert response.status_code == 400
    assert payload["success"] is False
    assert "Cliente / Customer" in payload["missing_fields"]


def test_ready_source_must_become_an_individual_draft_before_send_or_cancel(isolated_dirs):
    active_dir = isolated_dirs["active"]
    source_name = "2026_4573"
    copy_sample(active_dir, f"{source_name}.xlsx")
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()
    editing = acquire_editing(client, source_name)
    metadata = edit_metadata(editing)

    send_response = client.post(
        f"/api/file/{source_name}/send",
        json={"customer_name": "Cliente", "_edit": metadata},
    )
    cancel_response = client.post(
        f"/api/file/{source_name}/cancel",
        json={"_edit": metadata},
    )

    assert send_response.status_code == 409
    assert send_response.get_json()["code"] == "draft_required"
    assert cancel_response.status_code == 409
    assert cancel_response.get_json()["code"] == "draft_required"
    assert (active_dir / f"{source_name}.xlsx").exists()

def test_draft_does_not_require_all_fields(isolated_dirs):
    """O draft deve aceitar dados parciais sem validação de campos obrigatórios."""
    active_dir = isolated_dirs["active"]
    copy_sample(active_dir, "2026_4572.xlsx")

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()

    editing = acquire_editing(client, "2026_4572")
    response = client.post(
        "/api/file/2026_4572/draft",
        json={
            "intervention_report": "Apenas rascunho parcial",
            "_edit": edit_metadata(editing),
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["success"] is True
    assert payload["status"] == "in_progress"
    assert payload["created_copy"] is True
    assert (isolated_dirs["active"] / "2026_4572.xlsx").exists()
    assert (isolated_dirs["active"] / payload["file"] / f"{payload['file']}.xlsx").exists()
    assert (isolated_dirs["active"] / payload["file"] / f"{payload['file']}__documento.json").exists()


def test_two_technicians_create_distinct_drafts_from_the_same_source(isolated_dirs):
    active_dir = isolated_dirs["active"]
    source_name = "2026_4580"
    source_path = copy_sample(active_dir, f"{source_name}.xlsx")
    original_bytes = source_path.read_bytes()
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    first_client = app.test_client()
    second_client = app.test_client()
    first = acquire_editing(first_client, source_name, "tab-a")
    second = acquire_editing(second_client, source_name, "tab-b")

    assert first["document_id"] != second["document_id"]

    def save_draft(client, editing, client_id, report):
        return client.post(
            f"/api/file/{source_name}/draft",
            json={
                "intervention_report": report,
                "technician_records": [{"technician": "João Freire"}],
                "_edit": edit_metadata(editing, client_id=client_id),
            },
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        first_future = executor.submit(
            save_draft, first_client, first, "tab-a", "Rascunho exclusivo A"
        )
        second_future = executor.submit(
            save_draft, second_client, second, "tab-b", "Rascunho exclusivo B"
        )
        responses = [first_future.result(), second_future.result()]

    payloads = [response.get_json() for response in responses]
    assert [response.status_code for response in responses] == [200, 200]
    assert payloads[0]["file"] != payloads[1]["file"]
    assert payloads[0]["document_id"] != payloads[1]["document_id"]
    assert source_path.read_bytes() == original_bytes

    reports = set()
    for payload in payloads:
        bundle = active_dir / payload["file"]
        document_path = bundle / f"{payload['file']}__documento.json"
        reports.add(json.loads(document_path.read_text(encoding="utf-8"))["intervention_report"])
    assert reports == {"Rascunho exclusivo A", "Rascunho exclusivo B"}


def test_installation_form_places_work_number_between_store_and_contract(isolated_dirs):
    app = create_app(
        file_service=FileService(isolated_dirs["active"]),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )

    response = app.test_client().get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    contact_position = html.index('id="site-contact"')
    phone_position = html.index('id="site-phone"')
    store_position = html.index('id="local-store"')
    work_position = html.index('id="work-number"')
    contract_position = html.index('id="contract-number"')
    assert contact_position < phone_position < store_position < work_position < contract_position
    assert 'id="work-number" type="hidden" data-field="work_number"' in html
    assert 'id="work-number-link"' in html
    assert 'target="_blank"' in html
    assert 'rel="noopener noreferrer"' in html
    assert 'role="link"' in html
    assert 'aria-disabled="true"' in html


def test_document_preview_returns_html(isolated_dirs):
    active_dir = isolated_dirs["active"]
    copy_sample(active_dir, "2026_4572.xlsx")

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()

    response = client.post(
        "/api/file/2026_4572/document-preview",
        json={
            "service_number": "2026_4572",
            "customer_name": "Cliente Preview",
            "work_number": "0042",
            "intervention_report": "Relatório em HTML",
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["success"] is True
    assert "<!DOCTYPE html>" in payload["html"]
    assert "Cliente Preview" in payload["html"]
    assert "Folha de Serviço" in payload["html"]
    assert "Relatório de Intervenção" in payload["html"]
    assert "Dados do cliente" in payload["html"]
    assert "Dados da instalação" in payload["html"]
    assert "N.º de obra" in payload["html"]
    assert "0042" in payload["html"]
    assert "NIF / N.º de IVA" in payload["html"]
    assert "Trabalhos a efetuar" in payload["html"]
    assert "Descrição" in payload["html"]
    assert "Registo de técnicos e horas" in payload["html"]
    assert "Técnico responsável" in payload["html"]
    assert "Relatorio tecnico da intervencao" not in payload["html"]
    assert "data:image/png;base64," in payload["html"]


def test_file_not_found_returns_404(isolated_dirs):
    """Pedir um ficheiro inexistente deve devolver 404."""
    active_dir = isolated_dirs["active"]
    active_dir.mkdir(parents=True, exist_ok=True)

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )
    client = app.test_client()

    response = client.get("/api/file/inexistente")
    payload = response.get_json()

    assert response.status_code == 404
    assert payload["success"] is False


def test_web_editor_exposes_features_one_to_four(isolated_dirs):
    active_dir = isolated_dirs["active"]
    copy_sample(active_dir, "2026_4572.xlsx")
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )

    html = app.test_client().get("/?file=2026_4572").get_data(as_text=True)

    assert 'id="materials-panel"' in html
    assert 'id="materials-used-yes"' in html
    assert 'id="client-not-present"' in html
    total_marker = html.index('data-repeat-field="total_hours"')
    assert "readonly" not in html[total_marker - 160:total_marker + 220]

    positions = [
        html.index(f">{label}</span>")
        for label in ("SADI", "VSS", "SADCO", "SADIR", "SADEI", "SCA", "EAS", "SADG", "SCH", "OTHER")
    ]
    assert positions == sorted(positions)


def test_web_api_requires_signature_when_customer_is_present(isolated_dirs):
    active_dir = isolated_dirs["active"]
    draft_name = "2026_4572_2026-07-14_JF"
    copy_sample(active_dir, f"{draft_name}.xlsx")
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )

    client = app.test_client()
    editing = acquire_editing(client, draft_name)
    response = client.post(
        f"/api/file/{draft_name}/send",
        json={
            "customer_name": "Cliente presente",
            "_edit": edit_metadata(editing),
        },
    )
    payload = response.get_json()

    assert response.status_code == 400
    assert payload["success"] is False
    assert "Assinatura Cliente" in payload["missing_fields"]

def test_web_api_allows_client_absence_without_signature_and_records_it(isolated_dirs):
    active_dir = isolated_dirs["active"]
    draft_name = "2026_4572_2026-07-14_JF"
    copy_sample(active_dir, f"{draft_name}.xlsx")
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )

    client = app.test_client()
    editing = acquire_editing(client, draft_name)
    response = client.post(
        f"/api/file/{draft_name}/send",
        json={
            "customer_name": "Cliente ausente",
            "client_not_present": True,
            "customer_signature_date": "2026-07-10",
            "_edit": edit_metadata(editing),
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["success"] is True
    archived_dir = isolated_dirs["archived"] / draft_name
    document = json.loads(
        (archived_dir / f"{draft_name}__documento.json").read_text(encoding="utf-8")
    )
    html = (archived_dir / f"{draft_name}__folha_final.html").read_text(encoding="utf-8")

    assert document["client_not_present"] is True
    assert "Cliente não presente na obra" in html
    assert "Assinatura dispensada" in html
    assert not (archived_dir / f"{draft_name}__assinatura_cliente.png").exists()

def test_web_api_rejects_invalid_manual_total_on_send(isolated_dirs):
    active_dir = isolated_dirs["active"]
    draft_name = "2026_4572_2026-07-14_JF"
    copy_sample(active_dir, f"{draft_name}.xlsx")
    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
    )

    client = app.test_client()
    editing = acquire_editing(client, draft_name)
    response = client.post(
        f"/api/file/{draft_name}/send",
        json={
            "customer_name": "Cliente",
            "client_not_present": True,
            "technician_records": [
                {"technician": "João Freire", "total_hours": "oito horas"}
            ],
            "_edit": edit_metadata(editing),
        },
    )
    payload = response.get_json()

    assert response.status_code == 400
    assert payload["success"] is False
    assert payload["invalid_fields"] == ["Total de horas do técnico 1"]

def test_send_queues_customer_email_with_technician_cc_and_teams_notice(
    isolated_dirs,
    monkeypatch,
):
    import src.web.application as application_module

    monkeypatch.setattr(application_module, "ACTIVE_AUTH_PROVIDER", "microsoft")
    monkeypatch.setattr(application_module, "AUTH_ENABLED", True)
    active_dir = isolated_dirs["active"]
    draft_name = "2026_7001_2026-08-05_JF"
    copy_sample(active_dir, f"{draft_name}.xlsx")

    class DummyMicrosoftAuth:
        @staticmethod
        def user_from_session(payload):
            if not payload:
                return None
            return SimpleNamespace(
                id=payload["id"],
                display_name=payload["display_name"],
                email=payload["email"],
            )

    class RecordingQueue:
        def __init__(self):
            self.calls = []

        def enqueue(self, kind, payload, *, job_id=None):
            self.calls.append((kind, payload, job_id))
            return {"id": job_id, "status": "pending"}

    queue = RecordingQueue()
    mail_service = GraphMailService(
        GraphMailConfig(
            tenant_id="tenant",
            client_id="client",
            client_secret="secret",
            sender="service@sensorpoint.pt",
        )
    )

    class RecordingTeams:
        @staticmethod
        def prepare_service_sent_notification(*, service_number):
            return {
                "service_number": service_number,
                "payload": {"type": "message", "attachments": [{"content": {}}]},
            }

        @staticmethod
        def status():
            return {"enabled": True, "configured": True}

    app = create_app(
        file_service=FileService(active_dir),
        archive_service=ArchiveService(),
        mail_service=mail_service,
        microsoft_auth_service=DummyMicrosoftAuth(),
        editing_state_service=EditingStateService(isolated_dirs["editing"]),
        graph_sync_queue=queue,
        teams_notification_service=RecordingTeams(),
    )
    client = app.test_client()
    with client.session_transaction() as user_session:
        user_session["microsoft_user"] = {
            "id": "technician-id",
            "display_name": "Técnico Teste",
            "email": "tecnico@sensorpoint.pt",
        }

    editing = acquire_editing(client, draft_name, "mail-tab")
    response = client.post(
        f"/api/file/{draft_name}/send",
        json={
            "service_number": "2026_7001",
            "customer_name": "Cliente Email",
            "customer_email": "cliente@example.com",
            "client_not_present": True,
            "intervention_report": "Serviço concluído",
            "technician_records": [{"technician": "João Freire"}],
            "_edit": edit_metadata(editing, client_id="mail-tab"),
        },
    )

    payload = response.get_json()
    assert response.status_code == 200, payload
    assert payload["email_status"] == "pending"
    assert payload["email_recipient"] == "cliente@example.com"
    assert payload["teams_status"] == "pending"
    assert len(queue.calls) == 1
    kind, queued_payload, job_id = queue.calls[0]
    assert kind == "archive_and_mail_local"
    assert job_id.startswith("send:")
    assert queued_payload["mail"]["to"] == "cliente@example.com"
    assert queued_payload["mail"]["cc"] == ["tecnico@sensorpoint.pt"]
    assert queued_payload["teams"]["service_number"] == "2026_7001"

    archived_excel = Path(queued_payload["archived_path"])
    workbook = openpyxl.load_workbook(archived_excel)
    try:
        assert workbook["LINK"].sheet_state == "hidden"
        assert workbook.active.title == "FS"
    finally:
        workbook.close()
