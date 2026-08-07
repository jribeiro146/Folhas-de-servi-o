"""
Folhas de Serviço — Testes do Excel Service.

Usa uma cópia do ficheiro Excel de exemplo como fixture.
Testa leitura, escrita e validação sem tocar em ficheiros de produção.
"""

import shutil
import tempfile
import datetime as dt
from pathlib import Path

import pytest

# Resolve o caminho da fixture
FIXTURE_DIR = Path(__file__).parent / "fixtures"
SAMPLE_FILE = FIXTURE_DIR / "test_sample.xlsx"


def get_test_copy():
    """Cria uma cópia temporária do ficheiro de teste."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    shutil.copy2(str(SAMPLE_FILE), tmp.name)
    return Path(tmp.name)


def create_shifted_link_file():
    """Cria um workbook com as colunas essenciais do LINK deslocadas."""
    import openpyxl

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LINK"
    wb.create_sheet("FS")

    headers = [
        "Folha nº",
        "Pedido por:",
        "Email",
        "Telefone",
        "Contacto",
        "Telefone",
        "Contrato nº",
        "Loja nº",
        "Data\npedido",
        "Cliente nº",
        "ASSIST",
        "MAN",
        "COL.SERV",
        "GAR",
        "INST",
        "PIQ",
        "FORM",
        "REP.OF",
        "ACOMP.COM",
        "SADI",
        "CCTV",
        "PA/VA",
        "SAI",
        "EXT",
        "SCA",
        "EAS",
        "SADG",
        "SCH",
        "OTHER",
        "Cliente nome",
        "Local",
        "Morada",
        "Cód.\nPostal",
        "CP",
        "NIF",
        "Avaria reportada",
        "Fim",
        "Data serviço",
    ]
    values = [
        "2023/3021",
        "Sr. João Bernardo",
        "joao@acordo.pt",
        "917212167",
        "Sr. Ricardo António",
        "924222685",
        "n/a",
        "n/a",
        "22/11/2023",
        "1103",
        "X",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "X",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "ACORDO - Comércio e Serviços, Lda",
        "Hospital Misericórdia",
        "Av. Sanches de Miranda, 30",
        "7006-805 Évora",
        "70",
        "501398392",
        "Corrigir as ligações",
        dt.time(14, 30),
        dt.date(2026, 6, 16),
    ]

    for column_index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=column_index, value=header)
    for column_index, value in enumerate(values, start=1):
        ws.cell(row=3, column=column_index, value=value)

    wb.save(tmp.name)
    wb.close()
    return Path(tmp.name)


def create_work_number_link_file(header, value, position):
    """Cria um LINK mínimo com a coluna da obra numa posição variável."""
    import openpyxl

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "LINK"
    workbook.create_sheet("FS")

    headers = ["Folha nº", "Email", "Cliente nome", "Local", "NIF"]
    headers.insert(position, header)
    for column_index, current_header in enumerate(headers, start=1):
        worksheet.cell(row=2, column=column_index, value=current_header)
        worksheet.cell(
            row=3,
            column=column_index,
            value=value if column_index == position + 1 else "fixture",
        )

    workbook.save(tmp.name)
    workbook.close()
    return Path(tmp.name)


class TestExcelServiceValidation:
    """Testes de validação do ficheiro Excel."""

    def test_file_not_found(self):
        from src.services.excel_service import ExcelService
        with pytest.raises(FileNotFoundError):
            ExcelService("ficheiro_inexistente.xlsx")

    def test_valid_file(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        service.validate_sheets()  # Não deve lançar excepção

    def test_sheets_exist(self):
        import openpyxl
        wb = openpyxl.load_workbook(str(SAMPLE_FILE), read_only=True)
        assert "LINK" in wb.sheetnames
        assert "FS" in wb.sheetnames
        wb.close()


class TestExcelServiceRead:
    """Testes de leitura da sheet LINK."""

    def test_read_link_returns_dict(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        data = service.read_link()
        assert isinstance(data, dict)
        assert len(data) > 0

    def test_read_folha_numero(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        numero = service.get_folha_numero()
        assert numero is not None
        assert "2026_4572" in str(numero)

    def test_read_cliente_nome(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        data = service.read_link()
        # Coluna AE = Cliente nome
        assert data["AE"] is not None
        assert "CUF" in str(data["AE"])

    def test_read_checkbox_field(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        data = service.read_link()
        # Coluna M = MAN (marcado com "X" no ficheiro de exemplo)
        assert data["M"] == "X"

    def test_read_link_as_form_data(self):
        from src.services.excel_service import ExcelService
        service = ExcelService(SAMPLE_FILE)
        form_data = service.read_link_as_form_data()
        assert isinstance(form_data, dict)
        # Checkbox MAN deve ser True
        assert form_data["MAN"] is True
        # Checkbox ASSIST deve ser False (não marcado)
        assert form_data["ASSIST"] is False

    def test_read_link_uses_headers_when_columns_shift(self):
        from src.services.excel_service import ExcelService
        test_file = create_shifted_link_file()
        try:
            service = ExcelService(test_file)
            form_data = service.read_link_as_form_data()

            assert form_data["Folha nº"] == "2023/3021"
            assert form_data["Pedido por"] == "Sr. João Bernardo"
            assert form_data["Telefone"] == "917212167"
            assert form_data["Telefone (2)"] == "924222685"
            assert form_data["Cliente nome"] == "ACORDO - Comércio e Serviços, Lda"
            assert form_data["Local"] == "Hospital Misericórdia"
            assert form_data["Cód. Postal"] == "7006-805 Évora"
            assert form_data["ASSIST"] is True
            assert form_data["SADI"] is True
            assert form_data["Avaria reportada"] == "Corrigir as ligações"
            assert "Fim" not in form_data
            assert "Data serviço" not in form_data
        finally:
            test_file.unlink(missing_ok=True)

    @pytest.mark.parametrize(
        ("header", "position"),
        [
            ("N.º de obra", 0),
            ("Nº Obra", 2),
            ("Nº de Obra", 3),
            ("N.º obra", 4),
            ("Número de obra", 5),
        ],
    )
    def test_read_work_number_alias_in_variable_column(self, header, position):
        from src.document_schema import document_from_excel_and_extra
        from src.services.excel_service import ExcelService

        test_file = create_work_number_link_file(header, 22, position)
        try:
            form_data = ExcelService(test_file).read_link_as_form_data()
            document = document_from_excel_and_extra(form_data, {})

            assert form_data["N.º de obra"] == 22
            assert document["work_number"] == "0022"
        finally:
            test_file.unlink(missing_ok=True)

    def test_missing_work_number_header_returns_empty_value(self):
        from src.services.excel_service import ExcelService

        assert ExcelService(SAMPLE_FILE).read_link_as_form_data()["N.º de obra"] == ""

    def test_read_work_number_from_column_ah_preserves_shifted_postal_code(self):
        import openpyxl

        from src.services.excel_service import ExcelService

        test_file = get_test_copy()
        try:
            workbook = openpyxl.load_workbook(test_file)
            worksheet = workbook["LINK"]
            worksheet.insert_cols(34)
            worksheet.cell(row=2, column=34, value="Nº Obra")
            worksheet.cell(row=3, column=34, value=22)
            workbook.save(test_file)
            workbook.close()

            form_data = ExcelService(test_file).read_link_as_form_data()

            assert form_data["N.º de obra"] == 22
            assert form_data["Cód. Postal"] == "1998-018 Lisboa"
        finally:
            test_file.unlink(missing_ok=True)

    def test_work_number_is_read_only(self):
        from src.services.excel_service import ExcelService

        test_file = create_work_number_link_file("N.º de obra", "0022", 3)
        try:
            service = ExcelService(test_file)
            service.write_link_from_form({"N.º de obra": "0036"})
            assert service.read_link_as_form_data()["N.º de obra"] == "0022"
        finally:
            test_file.unlink(missing_ok=True)


class TestExcelServiceWrite:
    """Testes de escrita na sheet LINK."""

    def test_write_and_read_back(self):
        from src.services.excel_service import ExcelService
        test_file = get_test_copy()
        try:
            service = ExcelService(test_file)

            # Escrever um valor novo num campo presente no LINK
            service.write_link({"E": "Teste de escrita automática"})

            # Ler de volta
            data = service.read_link()
            assert data["E"] == "Teste de escrita automática"
        finally:
            test_file.unlink(missing_ok=True)

    def test_write_does_not_touch_readonly(self):
        from src.services.excel_service import ExcelService
        test_file = get_test_copy()
        try:
            service = ExcelService(test_file)

            # Coluna A (Folha nº) é read_only
            original = service.read_link()["A"]
            service.write_link({"A": "VALOR_ALTERADO"})

            # Deve manter o valor original
            data = service.read_link()
            assert data["A"] == original
        finally:
            test_file.unlink(missing_ok=True)

    def test_write_checkbox(self):
        from src.services.excel_service import ExcelService
        test_file = get_test_copy()
        try:
            service = ExcelService(test_file)

            # Marcar ASSIST (coluna L)
            service.write_link({"L": "X"})

            data = service.read_link()
            assert data["L"] == "X"
        finally:
            test_file.unlink(missing_ok=True)

    def test_write_link_uses_headers_when_columns_shift(self):
        import openpyxl

        from src.services.excel_service import ExcelService
        test_file = create_shifted_link_file()
        try:
            service = ExcelService(test_file)

            service.write_link_from_form({
                "Cliente nome": "Cliente atualizado",
                "Telefone (2)": "910000000",
                "ASSIST": False,
                "Avaria reportada": "Nova avaria",
            })

            form_data = service.read_link_as_form_data()
            assert form_data["Cliente nome"] == "Cliente atualizado"
            assert form_data["Telefone (2)"] == "910000000"
            assert form_data["ASSIST"] is False
            assert form_data["Avaria reportada"] == "Nova avaria"

            wb = openpyxl.load_workbook(str(test_file), read_only=True, data_only=True)
            try:
                ws = wb["LINK"]
                assert ws["AD3"].value == "Cliente atualizado"
                assert ws["AE3"].value == "Hospital Misericórdia"
            finally:
                wb.close()
        finally:
            test_file.unlink(missing_ok=True)


class TestFieldMap:
    """Testes de integridade do field_map."""

    def test_all_columns_unique(self):
        from src.field_map import FIELD_MAP
        columns = [f.column for f in FIELD_MAP]
        assert len(columns) == len(set(columns)), "Colunas duplicadas no field_map"

    def test_all_labels_unique(self):
        from src.field_map import FIELD_MAP
        labels = [f.label for f in FIELD_MAP]
        assert len(labels) == len(set(labels)), "Labels duplicados no field_map"

    def test_field_by_column_lookup(self):
        from src.field_map import FIELD_BY_COLUMN
        assert "A" in FIELD_BY_COLUMN
        assert FIELD_BY_COLUMN["A"].label == "Folha nº"

    def test_required_fields_exist(self):
        from src.field_map import REQUIRED_FIELDS
        assert len(REQUIRED_FIELDS) > 0
        labels = [f.label for f in REQUIRED_FIELDS]
        assert "Folha nº" in labels
        assert "Cliente nome" in labels
