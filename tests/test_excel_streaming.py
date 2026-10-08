"""Synthetic workbooks: bounded XML reads, mapping and explicit handle release."""
import zipfile

import openpyxl
import pytest

from src.field_map import FIELD_MAP
from src.services.excel_service import ExcelService


def workbook(path, *, width=100):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "LINK"
    book.create_sheet("FS")["A1"] = "=LINK!D3"
    labels = ["Folha nº", "Pedido por:", "Email", "Telefone", "Contacto", "Telefone", "Cliente nome", "Local"]
    for column, label in enumerate(labels, start=4):
        sheet.cell(2, column, label)
        sheet.cell(3, column, f"synthetic-{column}")
    sheet.cell(2, width, "unused")
    sheet.cell(3, width, "=1+2")
    sheet.cell(500, 1, "irrelevant trailing row")
    book.save(path)
    book.close()
    return labels


@pytest.mark.parametrize("width", [40, 200])
def test_reading_link_opens_worksheet_a_bounded_number_of_times_and_closes_handles(tmp_path, monkeypatch, width):
    path = tmp_path / "synthetic.xlsx"
    workbook(path, width=width)
    original = zipfile.ZipFile.open
    streams = []

    def track(archive, name, *args, **kwargs):
        stream = original(archive, name, *args, **kwargs)
        if name == "xl/worksheets/sheet1.xml":
            streams.append(stream)
        return stream
    monkeypatch.setattr(zipfile.ZipFile, "open", track)
    data = ExcelService(path).read_link()
    assert data["A"] == "synthetic-4"
    # One dimension scan while opening the workbook and one bounded value scan.
    assert len(streams) <= 2
    assert all(stream.closed for stream in streams)
    replacement = tmp_path / "renamed.xlsx"
    path.replace(replacement)
    replacement.unlink()


def test_read_and_write_preserve_shifted_headers_duplicates_and_formulas(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook(path)
    before = path.read_bytes()
    service = ExcelService(path)
    fields = tuple(field for field in FIELD_MAP if field.label in {"Folha nº", "Telefone", "Telefone (2)", "Email"})
    values = service.read_link(fields)
    by_label = {field.label: values[field.column] for field in fields}
    assert by_label == {"Folha nº": "synthetic-4", "Telefone": "synthetic-7",
                        "Telefone (2)": "synthetic-9", "Email": "synthetic-6"}
    assert path.read_bytes() == before
    email = next(field for field in fields if field.label == "Email")
    service.write_link({email.column: "synthetic@example.invalid"})
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        assert book["FS"]["A1"].value == "=LINK!D3"
        assert book["LINK"].cell(3, 100).value == "=1+2"
        assert book["LINK"]["F3"].value == "synthetic@example.invalid"
    finally:
        book.close()
    assert service.read_link(fields)[email.column] == "synthetic@example.invalid"
    path.unlink()


def test_sparse_link_retains_legacy_fixed_column_fallback(tmp_path):
    path = tmp_path / "sparse.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "LINK"
    book.create_sheet("FS")
    book["LINK"]["A3"] = "synthetic sparse"
    book.save(path)
    book.close()
    data = ExcelService(path).read_link()
    assert data["A"] == "synthetic sparse"
    assert all(value is None for column, value in data.items() if column != "A")
    path.unlink()
