"""SADI isolation, content binding, hierarchy and atomic finalization regressions."""
import base64
from copy import deepcopy
import io
import json
import re
import uuid
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
import pytest

from src import maintenance_schema as m
from src.document_schema import normalize_document_payload
from src.services.document_data_service import DocumentDataService
from src.services.local_pdf_service import LocalPdfService
from src.services.editing_state_service import EditingStateService
from src.services.file_service import FileService
import openpyxl
from openpyxl.utils.cell import column_index_from_string
from src.field_map import FIELD_MAP
import src.web.application as web
import src.services.archive_service as archive_module


def synthetic_workbook(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "LINK"
    book.create_sheet("FS")
    for field in FIELD_MAP:
        sheet.cell(2, column_index_from_string(field.column), field.label)
    sheet["A3"] = "2026_9900"
    book.save(path)
    book.close()
    return path

def signature_image():
    image = Image.new("RGBA", (300, 100), (255, 255, 255, 0))
    ImageDraw.Draw(image).line([(10, 60), (40, 15), (70, 80), (110, 25), (160, 60), (240, 30)], fill="black", width=3)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode()


def checklist_photo(size=(1200, 700), color="#dceaf0"):
    image = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(image)
    draw.rectangle((size[0]//4, size[1]//6, size[0]*3//4, size[1]*5//6), fill="#28526a")
    draw.text((30, 30), "DEMO - equipamento ficticio", fill="#163246", font_size=26)
    stream = io.BytesIO(); image.save(stream, format="JPEG")
    return {"id": uuid.uuid4().hex, "name": "equipamento-ficticio.jpg", "caption": "Equipamento de demonstração", "error": "",
            "image": "data:image/jpeg;base64," + base64.b64encode(stream.getvalue()).decode()}


def answers(questions):
    return {key: {"answer": "OK", "justification": ""} for key, _ in questions}


def complete_site(location="Edifício 1", conventional=2, addressable=0, repeater=1):
    site = {"id": uuid.uuid4().hex, "version": m.VERSION, "location": location,
            "date": "2026-09-17", "technician": "Técnico fictício", "period": "quarterly",
            "general": answers(m.GENERAL), "configuration": {}, "signatures": {},
            "peripherals": answers(m.PERIPHERALS), "trials": answers(m.TRIALS), "coverage_areas": "Piso 1 — dados fictícios"}
    for kind, count in (("conventional", conventional), ("addressable", addressable), ("repeater", repeater)):
        site["configuration"][kind] = bool(count)
        site[kind] = []
        for index in range(count):
            unit = {key: ("2" if field_type == "number" else f"Teste {index + 1}") for key, _, field_type in m.FIELDS[kind]}
            unit.update(id=uuid.uuid4().hex, checks=answers(m.DEFINITION[kind]))
            site[kind].append(unit)
    return m.normalize_sites([site])[0]


def complete_document():
    document = normalize_document_payload({"customer_name": "Cliente fictício SADI", "service_number": "2026_9900",
        "service_types": {"manutencao": True}, "equipments": {"sadi": True}, "client_not_present": True,
        "technician_records": [{"technician": "Técnico fictício", "date": "2026-09-17", "start_time": "09:00", "end_time": "17:00"}],
        "maintenance_checklists": [complete_site(), complete_site("Edifício 2", 0, 1, 0)]})
    return document


def sign_document(document, secret):
    for site in document["maintenance_checklists"]:
        for role in ("technician", "customer"):
            site["signatures"][role] = m.make_signature(secret, document, site, role, "Pessoa fictícia", "2026-09-17", signature_image())


def preview_data(html):
    return json.loads(re.search(r'<script type="application/json" id="preview-data">(.*?)</script>', html, re.S)[1])


def selected_preview(html):
    data = preview_data(html)
    return next(item["html"] for item in data["documents"] if item["id"] == data["selected"])


def test_excel_definition_and_shared_peripherals():
    assert [len(m.DEFINITION[key]) for key in ("general", "conventional", "addressable", "repeater", "peripherals", "trials")] == [2, 6, 7, 4, 8, 7]
    document = complete_document()
    assert m.document_errors(document, "secret", signatures=False) == []
    document["maintenance_checklists"][0]["trials"]["B122"]["answer"] = ""
    errors = m.document_errors(document, "secret", signatures=False)
    assert len(errors) == 1 and "Periféricos / Ensaios" in errors[0]


def test_nc_requires_individual_justification_and_is_not_incomplete_after_explanation():
    site = complete_site()
    site["conventional"][0]["checks"]["B44"] = {"answer": "NC", "justification": ""}
    assert any("justificação NC" in error for error in m.site_errors(site))
    site["conventional"][0]["checks"]["B44"]["justification"] = "Tecla avariada; substituição recomendada."
    assert m.site_errors(site) == []


@pytest.mark.parametrize("period", ["monthly", "quarterly", "half_yearly"])
def test_partial_period_requires_coverage_once_per_site(period):
    site = complete_site(); site["period"] = period
    site["coverage_areas"] = ""
    assert any("áreas testadas" in error for error in m.site_errors(site))
    site["coverage_percent"] = "0"
    assert m.site_errors(site) == []
    site["coverage_percent"] = "101"
    assert any("percentagem inválida" in error for error in m.site_errors(site))


def test_counts_ids_and_hidden_sections():
    site = complete_site()
    site["conventional"][0]["used"] = "3"
    assert any("superior ao total" in error for error in m.site_errors(site))
    site["configuration"]["conventional"] = False
    assert m.site_errors(site) == []
    assert len(site["conventional"]) == 2
    site["configuration"]["conventional"] = True
    site["conventional"][0]["used"] = "-1"
    assert m.site_errors(site)


@pytest.mark.parametrize("change", ["answer", "customer", "site_id", "enabled", "name", "image", "peripheral", "coverage", "repeater_location"])
def test_signatures_are_bound_to_content_and_identity(change):
    document = complete_document(); sign_document(document, "secret")
    site = document["maintenance_checklists"][0]
    assert m.document_errors(document, "secret") == []
    if change == "answer": site["general"]["B22"]["answer"] = "NA"
    if change == "customer": document["customer_name"] = "Outro cliente"
    if change == "site_id": site["id"] = uuid.uuid4().hex
    if change == "enabled": site["configuration"]["repeater"] = False
    if change == "name": site["signatures"]["customer"]["name"] = "Outro nome"
    if change == "image": site["signatures"]["customer"]["image"] = "data:image/png;base64,eA=="
    if change == "peripheral": site["peripherals"]["B107"]["answer"] = "NA"
    if change == "coverage": site["coverage_areas"] = "Outra área"
    if change == "repeater_location": site["repeater"][0]["location"] = "Receção"
    assert not m.signature_valid("secret", document, site, "customer")
    m.clear_invalid_signatures(document, "secret")
    assert "customer" not in site["signatures"]
    if change != "customer": assert m.signature_valid("secret", document, document["maintenance_checklists"][1], "customer")


def test_roundtrip_and_non_maintenance_compatibility():
    document = complete_document(); sign_document(document, "secret")
    assert normalize_document_payload(document) == document
    assert m.document_errors(normalize_document_payload(document), "secret") == []
    document["service_types"]["manutencao"] = False
    assert m.document_errors(document, "secret") == []
    assert normalize_document_payload({})["maintenance_checklists"] == []


@pytest.mark.parametrize("change", ["add", "remove", "caption", "image", "notes"])
def test_photos_and_final_notes_invalidate_only_their_site(change):
    document = complete_document(); site = document["maintenance_checklists"][0]
    site["photos"] = [checklist_photo()]
    site["final_observations"] = "Conclusão da manutenção."
    sign_document(document, "secret")
    assert normalize_document_payload(document) == document
    if change == "add": site["photos"].append(checklist_photo())
    if change == "remove": site["photos"] = []
    if change == "caption": site["photos"][0]["caption"] = "Outra legenda"
    if change == "image": site["photos"][0]["image"] = checklist_photo(color="#ffcc00")["image"]
    if change == "notes": site["final_observations"] = "Conclusão alterada."
    m.clear_invalid_signatures(document, "secret")
    assert site["signatures"] == {}
    assert all(m.signature_valid("secret", document, document["maintenance_checklists"][1], role) for role in ("technician", "customer"))


def test_optional_empty_fields_preserve_existing_sadi2_signatures():
    document = complete_document()
    for site in document["maintenance_checklists"]:
        site.pop("photos"); site.pop("final_observations")
        site.pop("signature_drafts")
    sign_document(document, "secret")
    loaded = normalize_document_payload(document)
    assert m.document_errors(loaded, "secret") == []


def test_unsigned_capture_survives_drafts_but_does_not_waive_client_signature(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document(); sign_document(document, app.secret_key)
    site = document["maintenance_checklists"][0]
    del site["signatures"]["customer"]
    site["signature_drafts"]["customer"] = {"name": "Cliente fictício", "date": "2026-09-21", "image": signature_image()}
    before = deepcopy(site)
    result = client.post(f"/api/file/{draft.stem}/draft", json={**document, "_edit": metadata})
    assert result.status_code == 200, result.get_json()
    loaded = client.get(f"/api/file/{draft.stem}").get_json()["document"]
    assert loaded["maintenance_checklists"][0] == before
    assert m.signature_valid(app.secret_key, loaded, loaded["maintenance_checklists"][0], "technician")
    assert any("assinatura do cliente" in error for error in m.document_errors(loaded, app.secret_key))
    html = client.post(f"/api/file/{draft.stem}/document-preview", json={**loaded,"_maintenance_site_id":site["id"],"_auto_print":True}).get_json()["html"]
    assert "Cliente fictício</strong>" not in html and html.count('class="signature-image has-signature"') == 1
    # The service-sheet exception never dispenses with the checklist's customer signature.
    assert loaded["client_not_present"] is True
    assert client.post(f"/api/file/{draft.stem}/send", json={**loaded,"_edit":{**metadata,"idempotency_key":uuid.uuid4().hex}}).status_code == 400
    assert not renders


def test_unsigned_capture_rejects_external_or_invalid_images():
    site = complete_site()
    site["signature_drafts"] = {"customer":{"name":"Teste","date":"","image":"https://example.invalid/signature.png"}}
    normalized = m.normalize_sites([site])[0]
    assert normalized["signature_drafts"]["customer"]["image"] == ""


@pytest.mark.parametrize("bad", ["https://example.invalid/photo.jpg", "data:image/svg+xml;base64,PHN2Zz4=", "data:image/jpeg;base64,eA=="])
def test_untrusted_photos_are_not_renderable_and_block_completion(bad):
    site = complete_site(); photo = checklist_photo(); photo["image"] = bad; site["photos"] = [photo]
    normalized = m.normalize_sites([site])[0]
    assert normalized["photos"][0]["image"] == ""
    assert normalized["photos"][0]["error"]
    assert m.site_errors(normalized)
    assert m.normalize_sites([normalized]) == [normalized]


def test_photo_dimensions_count_and_duplicates_are_validated():
    site = complete_site(); site["photos"] = [checklist_photo(size=(1601, 500))]
    assert m.site_errors(site)
    photo = checklist_photo(); site["photos"] = [photo, deepcopy(photo)]
    assert any("repetido" in error for error in m.site_errors(site))
    site["photos"] = [checklist_photo() for _ in range(11)]
    assert any("máximo" in error for error in m.site_errors(site))


def legacy_site():
    site = complete_site()
    site["version"] = "sadi-1"
    old = {key: site.pop(key) for key in ("peripherals", "trials") + m.PERIPHERAL_FIELDS}
    for unit in site["conventional"]:
        unit.update(deepcopy(old))
    site["repeater"][0]["location"] = ""
    site["signatures"] = {"technician": {"token": "old"}, "customer": {"token": "old"}}
    return site


def test_v1_migration_preserves_originals_and_only_copies_unanimous_answers():
    old = legacy_site()
    before = deepcopy(old)
    site = m.normalize_sites([old])[0]
    assert old == before
    assert site["version"] == m.VERSION and site["signatures"] == {}
    assert site["peripherals"] == old["conventional"][0]["peripherals"]
    assert site["trials"] == old["conventional"][0]["trials"]
    assert len(site["peripheral_history"]) == 2
    assert "peripherals" not in site["conventional"][0]
    assert site["conventional"][0]["id"] == old["conventional"][0]["id"]
    assert m.normalize_sites([site]) == [site]
    assert m.site_errors(site) == ["Repetidor 1: Local do repetidor"]


def test_migration_preserves_conflicting_and_inactive_records_without_guessing():
    old = legacy_site()
    old["conventional"][0]["peripherals"]["B107"] = {"answer": "NC", "justification": "Fonte avariada"}
    old["conventional"][0].update(coverage_percent="25", coverage_areas="Piso 1")
    old["conventional"][1].update(coverage_percent="50", coverage_areas="Piso 2")
    old["addressable"] = [deepcopy(old["conventional"][0])]
    old["addressable"][0]["trials"]["B121"]["answer"] = "NA"
    site = m.normalize_sites([old])[0]
    assert site["peripherals"]["B107"] == {"answer": "", "justification": ""}
    assert site["trials"]["B121"]["answer"] == "OK"
    assert site["coverage_percent"] == site["coverage_areas"] == ""
    assert len(site["peripheral_history"]) == 3
    assert site["peripheral_history"][0]["peripherals"]["B107"]["justification"] == "Fonte avariada"
    assert site["peripheral_history"][2]["active"] is False
    assert m.normalize_sites([site]) == [site]


def test_peripherals_are_always_required_and_survive_equipment_changes():
    site = complete_site(conventional=0, addressable=0, repeater=0)
    assert m.site_errors(site) == []
    site["peripherals"]["B107"] = {"answer": "NC", "justification": ""}
    assert any("justificação NC" in error for error in m.site_errors(site))
    site["peripherals"]["B107"]["justification"] = "Fonte por substituir"
    assert m.site_errors(site) == []
    site["trials"]["B122"]["answer"] = ""
    assert len(m.site_errors(site)) == 1
    assert m.normalize_sites([site])[0]["peripherals"] == site["peripherals"]


@pytest.fixture
def demo(monkeypatch):
    tmp_path = Path(tempfile.mkdtemp(prefix="sadi-test-"))
    monkeypatch.setattr(web, "STORAGE_BACKEND", "local")
    monkeypatch.setattr(web, "MAIL_ENABLED", False)
    monkeypatch.setattr(web, "TEAMS_NOTIFICATIONS_ENABLED", False)
    monkeypatch.setattr(web, "APP_DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(archive_module, "EXCEL_ARQUIVADAS_DIR", tmp_path / "archive")
    source = synthetic_workbook(tmp_path / "active" / "2026_9900.xlsx")
    files = FileService(source.parent)
    draft = files.create_draft_copy(source, "Técnico fictício")
    renders = []
    def render(html, destination):
        renders.append(html)
        destination.write_bytes(b"%PDF-1.4\nsynthetic test renderer")
    app = web.create_app(file_service=files, editing_state_service=EditingStateService(tmp_path / "editing"), local_pdf_service=LocalPdfService(renderer=render))
    client = app.test_client()
    editing = client.post(f"/api/file/{draft.stem}/lease", json={"client_id": "qa"}).get_json()["editing"]
    metadata = {"document_id": editing["document_id"], "client_id": "qa", "lease_token": editing["lease"]["token"], "base_revision": editing["revision"], "idempotency_key": uuid.uuid4().hex}
    return app, client, draft, metadata, renders, tmp_path


def test_api_sign_draft_reopen_and_invalidate(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    document["maintenance_checklists"][0]["photos"] = [checklist_photo()]
    document["maintenance_checklists"][0]["final_observations"] = "Primeira linha\nSegunda linha — verificação concluída."
    for site in document["maintenance_checklists"]:
        for role in ("technician", "customer"):
            result = client.post(f"/api/file/{draft.stem}/maintenance/sign", json={"_edit": metadata, "document": document, "site_id": site["id"], "role": role, "name": "Pessoa fictícia", "date": "2026-09-17", "image": signature_image()})
            assert result.status_code == 200, result.get_json()
            site["signatures"][role] = result.get_json()["signature"]
    result = client.post(f"/api/file/{draft.stem}/draft", json={**document, "_edit": metadata})
    assert result.status_code == 200, result.get_json()
    persisted = DocumentDataService(draft).read()
    assert m.document_errors(persisted, app.secret_key) == []
    loaded = client.get(f"/api/file/{draft.stem}").get_json()["document"]
    assert loaded["maintenance_checklists"] == document["maintenance_checklists"]
    document["maintenance_checklists"][0]["general"]["B22"]["answer"] = "NA"
    result = client.post(f"/api/file/{draft.stem}/maintenance/validate", json={"document": document}).get_json()
    assert result["sites"][0]["signed"] == {"technician": False, "customer": False}
    assert result["sites"][1]["signed"] == {"technician": True, "customer": True}


def test_api_autosave_photos_and_reject_invalid_at_finalization(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document(); site = document["maintenance_checklists"][0]
    site["photos"] = [checklist_photo()]; site["final_observations"] = "Nota final de teste"
    result = client.post(f"/api/file/{draft.stem}/autosave", json={"document": document, "_edit": metadata})
    assert result.status_code == 200
    reopened = client.get(f"/api/file/{draft.stem}").get_json()
    assert "Nota final de teste" in json.dumps(reopened, ensure_ascii=False)
    site["photos"][0]["image"] = "https://example.invalid/must-not-fetch.jpg"
    result = client.post(f"/api/file/{draft.stem}/maintenance/sign", json={"_edit": metadata, "document": document,
        "site_id": site["id"], "role": "technician", "name": "Pessoa fictícia", "date": "2026-09-17", "image": signature_image()})
    assert result.status_code == 200
    result = client.post(f"/api/file/{draft.stem}/send", json={**document, "_edit": metadata})
    assert result.status_code == 400 and not renders and draft.exists()
    result = client.post(f"/api/file/{draft.stem}/document-preview", json={**document, "_maintenance_site_id": site["id"], "_auto_print": True})
    assert result.status_code == 200 and "must-not-fetch" not in result.get_json()["html"]


def test_finalization_all_pdfs_and_replay(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    document["maintenance_checklists"][0]["photos"] = [checklist_photo()]
    document["maintenance_checklists"][0]["final_observations"] = "Observação final fictícia."
    sign_document(document, app.secret_key)
    payload = {**document, "_edit": metadata}
    result = client.post(f"/api/file/{draft.stem}/send", json=payload)
    assert result.status_code == 200, result.get_json()
    response = result.get_json()
    assert len(renders) == 3 and not draft.exists()
    page = client.get(response["maintenance_bundle_url"])
    assert page.status_code == 200 and "3 PDFs" in page.text
    simulation = client.post(response["maintenance_bundle_url"] + "/simulate").get_json()
    assert simulation["simulated"]
    assert len(simulation["attachments"]) == 1
    assert len(simulation["stored_only"]) == 2
    assert "__SADI_" not in simulation["attachments"][0]
    assert all("__SADI_" in name for name in simulation["stored_only"])
    assert "Apenas a folha de serviço segue para o cliente" in page.text
    archive = next((root / "archive").glob("*/maintenance-manifest.json")).parent
    manifest = archive / "maintenance-manifest.json"
    entries = json.loads(manifest.read_text(encoding="utf-8"))
    assert [entry["delivery"] for entry in entries] == ["customer_email", "archive_only", "archive_only"]
    for entry in entries:
        entry.pop("delivery")
    manifest.write_text(json.dumps(entries), encoding="utf-8")
    legacy_simulation = client.post(response["maintenance_bundle_url"] + "/simulate").get_json()
    assert legacy_simulation["attachments"] == simulation["attachments"]
    assert legacy_simulation["stored_only"] == simulation["stored_only"]
    htmls = sorted(archive.glob("*__SADI_*.html"))
    assert len(htmls) == 2
    html = htmls[0].read_text(encoding="utf-8")
    assert html.index("Fotografias") < html.index("Observações finais") < html.index("Assinatura do técnico")
    assert "Hora de saída" not in html and "Responsável SCIE" in html
    assert document["maintenance_checklists"][0]["photos"][0]["image"] in html
    assert "Observação final fictícia." in html and "Observação final fictícia." not in htmls[1].read_text(encoding="utf-8")
    assert "Central convencional 2" in html
    assert html.index("1. Sistema") < html.index("2. Centrais") < html.index("3. Repetidores") < html.index("4. Periféricos e ensaios")
    assert html.index("Local do repetidor") < html.index("4. Periféricos e ensaios")
    assert html.count("4. Periféricos e ensaios") == 1
    for _, label in m.PERIPHERALS + m.TRIALS:
        assert html.count(label.removesuffix(' (*)')) == 1
    assert client.post(f"/api/file/{draft.stem}/send", json=payload).get_json() == response
    assert len(renders) == 3
    assert client.get(response["maintenance_bundle_url"] + "/pdf/maintenance-manifest.json").status_code == 404


def test_pdf_failure_retains_draft_and_retry_builds_complete_bundle(demo, monkeypatch):
    app, client, draft, metadata, renders, root = demo
    document = complete_document(); sign_document(document, app.secret_key)
    original = LocalPdfService.export_html_pdf
    attempts = []
    def failing(self, html, destination):
        attempts.append(html)
        if len(attempts) == 2: raise RuntimeError("synthetic renderer failure")
        return original(self, html, destination)
    monkeypatch.setattr(LocalPdfService, "export_html_pdf", failing)
    payload = {**document, "_edit": metadata}
    result = client.post(f"/api/file/{draft.stem}/send", json=payload)
    assert result.status_code == 500 and draft.exists()
    assert not list((root / "archive").glob("*/maintenance-manifest.json"))
    result = client.post(f"/api/file/{draft.stem}/send", json=payload)
    assert result.status_code == 200, result.get_json()
    assert len(list((root / "archive").glob("*/*.pdf"))) == 3


def test_api_shared_checks_and_repeater_location_are_required_for_completion(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    site = complete_site(conventional=0, addressable=0, repeater=1)
    document["maintenance_checklists"] = [site]
    site["peripherals"]["B107"] = {"answer": "NC", "justification": ""}
    site["repeater"][0]["location"] = ""
    def sign():
        return client.post(f"/api/file/{draft.stem}/maintenance/sign", json={
            "_edit": metadata, "document": document, "site_id": site["id"], "role": "technician",
            "name": "Técnico fictício", "date": "2026-09-17", "image": signature_image()})
    assert sign().status_code == 200
    errors = client.post(f"/api/file/{draft.stem}/maintenance/validate", json={"document": document}).get_json()["sites"][0]["errors"]
    assert len(errors) == 2 and any("Local do repetidor" in error for error in errors)
    site["repeater"][0]["location"] = "Receção"
    site["peripherals"]["B107"]["justification"] = "Fonte por substituir"
    assert sign().status_code == 200
    # No equipment toggle removes the permanent fourth section or its requirements.
    site["configuration"]["repeater"] = False
    site["trials"]["B122"]["answer"] = ""
    assert sign().status_code == 200
    assert client.post(f"/api/file/{draft.stem}/maintenance/validate", json={"document": document}).get_json()["sites"][0]["errors"]
    html = client.post(f"/api/file/{draft.stem}/document-preview", json={
        **document, "_maintenance_site_id": site["id"], "_auto_print": True}).get_json()["html"]
    assert html.index("3. Repetidores") < html.index("4. Periféricos e ensaios")
    assert "Por responder" in html and "Repetidor 1" not in html
    assert all(html.count(label.removesuffix(" (*)")) == 1 for _, label in m.PERIPHERALS + m.TRIALS)


def test_migrated_draft_save_reopen_preserves_answers_and_history(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    document["maintenance_checklists"] = [legacy_site()]
    response = client.post(f"/api/file/{draft.stem}/draft", json={**document, "_edit": metadata})
    assert response.status_code == 200, response.get_json()
    saved = DocumentDataService(draft).read()["maintenance_checklists"][0]
    reopened = client.get(f"/api/file/{draft.stem}").get_json()["document"]["maintenance_checklists"][0]
    assert saved == reopened and saved["version"] == m.VERSION
    assert len(saved["peripheral_history"]) == 2 and saved["signatures"] == {}
    assert saved["peripherals"]["B107"]["answer"] == "OK"
    assert m.normalize_sites([saved]) == [saved]
    assert not renders


def test_api_rejects_incomplete_before_any_render_and_blocks_outside_demo(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    result = client.post(f"/api/file/{draft.stem}/send", json={**document, "_edit": metadata})
    assert result.status_code == 400 and not renders and draft.exists()
    site = document["maintenance_checklists"][0]
    site["general"]["B22"]["answer"] = ""
    result = client.post(f"/api/file/{draft.stem}/maintenance/sign", json={"_edit": metadata, "document": document, "site_id": site["id"], "role": "customer", "name": "Pessoa fictícia", "date": "2026-09-17", "image": signature_image()})
    assert result.status_code == 200
    app.config["MAINTENANCE_DEMO"] = False
    assert client.post(f"/api/file/{draft.stem}/maintenance/sign", json={}).status_code == 404


@pytest.mark.parametrize("pending", ["empty", "nc", "equipment", "coverage", "invalid_range"])
def test_incomplete_checklist_signs_reopens_but_cannot_finalize(demo, pending):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    sign_document(document, app.secret_key)
    site = document["maintenance_checklists"][0]
    if pending == "empty":
        site = m.normalize_sites([{"id": site["id"], "version": m.VERSION}])[0]
        document["maintenance_checklists"][0] = site
    elif pending == "nc":
        site["general"]["B22"] = {"answer": "NC", "justification": ""}
    elif pending == "equipment":
        site["repeater"][0]["location"] = ""
        site["conventional"][0]["brand"] = ""
    elif pending == "coverage":
        site["coverage_areas"] = ""
    elif pending == "invalid_range":
        site["conventional"][0]["used"] = "999"
    for role in ("technician", "customer"):
        result = client.post(f"/api/file/{draft.stem}/maintenance/sign", json={"_edit": metadata,
            "document": document, "site_id": site["id"], "role": role,
            "name": "Pessoa fictícia", "date": "2026-09-21", "image": signature_image()})
        assert result.status_code == 200, result.get_json()
        site["signatures"][role] = result.get_json()["signature"]
    checked = client.post(f"/api/file/{draft.stem}/maintenance/validate", json={"document": document}).get_json()
    assert checked["sites"][0]["errors"]
    assert checked["sites"][0]["signed"] == {"technician": True, "customer": True}
    saved = client.post(f"/api/file/{draft.stem}/draft", json={**document, "_edit": metadata})
    assert saved.status_code == 200, saved.get_json()
    loaded = client.get(f"/api/file/{draft.stem}").get_json()["document"]
    restored_site = loaded["maintenance_checklists"][0]
    assert restored_site == site
    assert all(m.signature_valid(app.secret_key, loaded, restored_site, role) for role in ("technician", "customer"))
    editing = client.post(f"/api/file/{draft.stem}/lease", json={"client_id": "qa"}).get_json()["editing"]
    metadata = {**metadata, "base_revision": editing["revision"], "lease_token": editing["lease"]["token"], "idempotency_key": uuid.uuid4().hex}
    finalized = client.post(f"/api/file/{draft.stem}/send", json={**loaded, "_edit": metadata})
    assert finalized.status_code == 400 and finalized.get_json()["missing_fields"]
    assert not renders and draft.exists()
    restored_site["final_observations"] = "Alteração posterior à assinatura"
    checked = client.post(f"/api/file/{draft.stem}/maintenance/validate", json={"document": loaded}).get_json()
    assert checked["sites"][0]["signed"] == {"technician": False, "customer": False}
    assert checked["sites"][1]["signed"] == {"technician": True, "customer": True}


@pytest.mark.parametrize("field,value", [("name", ""), ("date", "2026-02-30"), ("image", "data:image/png;base64,eA=="), ("role", "other")])
def test_incomplete_checklist_still_requires_a_valid_signature(demo, field, value):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    site = document["maintenance_checklists"][0]
    site["date"] = ""
    payload = {"_edit": metadata, "document": document, "site_id": site["id"], "role": "technician",
               "name": "Pessoa fictícia", "date": "2026-09-21", "image": signature_image()}
    payload[field] = value
    assert client.post(f"/api/file/{draft.stem}/maintenance/sign", json=payload).status_code == 400


def test_preview_selected_local_without_signing_or_saving(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    site = document["maintenance_checklists"][0]
    site["general"]["B22"] = {"answer": "NC", "justification": "Motivo ainda por guardar <script>"}
    before = DocumentDataService(draft).read()
    original = draft.read_bytes()
    for index, expected, excluded in ((0, "Central convencional 2", "Central endereçável 1"),
                                       (1, "Central endereçável 1", "Central convencional 1")):
        result = client.post(f"/api/file/{draft.stem}/document-preview", json={
            **document, "_maintenance_site_id": document["maintenance_checklists"][index]["id"],
            "_auto_print": True,
        })
        assert result.status_code == 200, result.get_json()
        html = result.get_json()["html"]
        assert expected in html and excluded not in html
        assert "RASCUNHO" not in html and html.count("Assinatura por recolher") == 2
        assert "Checklist por finalizar. Os campos e assinaturas em falta permanecem pendentes." not in html
        assert 'window.addEventListener("load"' in html
        assert document["maintenance_checklists"][1-index]["location"] not in html
        if index == 0:
            assert "Motivo ainda por guardar &lt;script&gt;" in html
    assert draft.read_bytes() == original
    assert DocumentDataService(draft).read() == before
    assert not renders and not list((root / "archive").glob("*/*.pdf"))
    # Previewing does not waive the finalization requirements.
    assert client.post(f"/api/file/{draft.stem}/send", json={**document, "_edit": metadata}).status_code == 400
    sheet = selected_preview(client.post(f"/api/file/{draft.stem}/document-preview", json=document).get_json()["html"])
    assert "Manutenção preventiva — SADI" not in sheet


def test_preview_empty_local_and_unsigned_equipment(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    site_id = document["maintenance_checklists"][0]["id"]
    document["maintenance_checklists"] = [{"id": site_id, "configuration": {"conventional": True},
                                          "conventional": [{"id": uuid.uuid4().hex}]}]
    result = client.post(f"/api/file/{draft.stem}/document-preview", json={**document, "_maintenance_site_id": site_id})
    assert result.status_code == 200, result.get_json()
    html = selected_preview(result.get_json()["html"])
    assert "Por responder" in html and "Por preencher" in html and "configuração por preencher" in html
    assert 'window.addEventListener("load"' not in html


def test_preview_uses_only_valid_signatures_and_active_equipment(demo):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    sign_document(document, app.secret_key)
    site = document["maintenance_checklists"][0]
    def preview():
        return selected_preview(client.post(f"/api/file/{draft.stem}/document-preview", json={
            **document, "_maintenance_site_id": site["id"]}).get_json()["html"])
    assert preview().count('class="signature-image has-signature"') == 2
    site["configuration"]["conventional"] = False
    html = preview()
    assert "Central convencional 1" not in html
    assert "Repetidor 1" in html
    assert 'class="signature-image has-signature"' not in html
    assert html.count("Assinatura por recolher") == 2


@pytest.mark.parametrize("selected", [None, 0, 1])
def test_preview_includes_sheet_and_every_local_with_independent_export(demo, selected):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    document["maintenance_checklists"][0]["observations"] = '</script><script>alert("example")</script>'
    payload = dict(document)
    if selected is not None:
        payload["_maintenance_site_id"] = document["maintenance_checklists"][selected]["id"]
    result = client.post(f"/api/file/{draft.stem}/document-preview", json=payload)
    assert result.status_code == 200, result.get_json()
    data = preview_data(result.get_json()["html"])
    assert data["selected"] == ("service" if selected is None else f"sadi-{selected}")
    assert [item["label"] for item in data["documents"]] == ["Folha de serviço", "SADI — Edifício 1", "SADI — Edifício 2"]
    sheet, first, second = [item["html"] for item in data["documents"]]
    assert "Relatório de Intervenção" in sheet and "Manutenção preventiva — SADI" not in sheet
    assert "Central convencional 2" in first and "Edifício 2" not in first
    assert "Central endereçável 1" in second and "Edifício 1" not in second
    assert '&lt;/script&gt;&lt;script&gt;alert(' in first
    # Every preview embeds the very same stylesheet and logo as the service sheet.
    styles, logo = web.load_document_assets(app.static_folder)
    for html in (sheet, first, second):
        assert styles in html and logo in html
        assert 'class="sheet-header"' in html and 'class="sheet-id-card"' in html
    exported = client.post(f"/api/file/{draft.stem}/document-preview", json={**payload, "_auto_print": True}).get_json()["html"]
    assert 'id="preview-data"' not in exported
    assert ("Relatório de Intervenção" in exported) == (selected is None)
    assert not renders and draft.exists()


@pytest.mark.parametrize("case", ["no_site", "duplicate", "inactive", "not_demo", "not_draft"])
def test_preview_rejects_invalid_target(demo, case):
    app, client, draft, metadata, renders, root = demo
    document = complete_document()
    site = document["maintenance_checklists"][0]
    target = site["id"]
    name = draft.stem
    expected = 400
    if case == "no_site": target = "does-not-exist"
    if case == "duplicate": document["maintenance_checklists"].append(site)
    if case == "inactive": document["equipments"]["sadi"] = False
    if case == "not_demo": app.config["MAINTENANCE_DEMO"] = False; expected = 404
    if case == "not_draft": name = "2026_9900"; expected = 409
    result = client.post(f"/api/file/{name}/document-preview", json={**document, "_maintenance_site_id": target})
    assert result.status_code == expected, result.get_json()
    assert not renders
